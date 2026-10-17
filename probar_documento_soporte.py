"""Prueba del documento soporte (Res. DIAN 167/2021): base soporte.db (ADR-001), numeración reservada (ADR-003)."""
import sqlite3, tempfile, threading
from pathlib import Path
from datetime import date
from parser_module.contable import soporte_db as sdb, proveedores as prov, soporte as sop
ok = True
def chk(c, m):
    global ok; ok &= bool(c); print(("✓ " if c else "✗ ") + m)

tmp = Path(tempfile.mkdtemp())
RUTA = tmp / "soporte.db"

# ── T7 · esquema, migraciones y semilla ─────────────────────────────────────
con = sdb.abrir(RUTA)
ver = con.execute("PRAGMA user_version").fetchone()[0]
chk(ver == len(sdb.MIGRACIONES), f"user_version = {ver} (migraciones aplicadas: {len(sdb.MIGRACIONES)})")
tablas = {r[0] for r in con.execute("SELECT name FROM sqlite_master WHERE type='table'")}
chk({"proveedor", "alias", "rango", "documento", "linea"} <= tablas, f"tablas de la sección 4 de 03: {sorted(tablas)}")
rg = [tuple(r) for r in con.execute("SELECT clase, prefijo, desde, hasta, ultimo, activo FROM rango")]
chk(rg == [("DS", "CTG", 1, 10000, 0, 1)], f"semilla: rango DS CTG 1–10000 activo, sin usar: {rg}")
con.close()
con = sdb.abrir(RUTA)   # abrir otra vez no repite migraciones ni semilla
chk(con.execute("SELECT count(*) FROM rango").fetchone()[0] == 1 and con.execute("PRAGMA user_version").fetchone()[0] == ver,
    "abrir dos veces no duplica la semilla ni cambia la versión")
chk(con.execute("PRAGMA foreign_keys").fetchone()[0] == 1, "las llaves foráneas están activas")

# ── reservar_numero ─────────────────────────────────────────────────────────
try: sdb.reservar_numero(con, "DS"); fuera = False
except RuntimeError: fuera = True
chk(fuera, "reservar_numero fuera de una transacción → error (debe ir en BEGIN IMMEDIATE)")
with sdb.transaccion(con):
    n1 = sdb.reservar_numero(con, "DS"); n2 = sdb.reservar_numero(con, "DS")
chk((n1, n2) == (("CTG", 1), ("CTG", 2)), f"números consecutivos: {n1}, {n2}")
try:
    with sdb.transaccion(con):
        sdb.reservar_numero(con, "DS"); raise ValueError("falla a mitad")
except ValueError: pass
chk(con.execute("SELECT ultimo FROM rango WHERE clase='DS'").fetchone()[0] == 2, "si la transacción falla, el contador vuelve atrás")

# ── ADR-001 · UNIQUE(prefijo, numero) ───────────────────────────────────────
pid = con.execute("INSERT INTO proveedor(tipo_id, num_id, nombre, creado) VALUES ('CC', '1234567', 'Doña Rosa', '2026-10-17T00:00:00Z')").lastrowid
def doc(numero, estado="generado"):
    return con.execute("INSERT INTO documento(clase, prefijo, numero, proveedor_id, semana_iso, fecha_generacion, total, estado, creado) "
                       "VALUES ('DS', 'CTG', ?, ?, '2026-W42', '2026-10-17', 50000, ?, '2026-10-17T00:00:00Z')", (numero, pid, estado)).lastrowid
d1 = doc(1)
try: doc(1); repetido = False
except sqlite3.IntegrityError: repetido = True
chk(repetido, "ADR-001: UNIQUE(prefijo,numero) ausente" if not repetido else "ADR-001: un número repetido → IntegrityError")
try: doc(3, "borrador"); raro = False
except sqlite3.IntegrityError: raro = True
chk(raro, "estado fuera de {generado, transmitido, anulado} → IntegrityError")
otros = [p for p in Path(".").rglob("*.py") if p.name != "soporte_db.py" and "venv" not in p.parts and p.name != Path(__file__).name
         and "sqlite3.connect" in p.read_text(encoding="utf8", errors="ignore")]
chk(not otros, "ADR-001: solo soporte_db.py abre soporte.db" + (f" (también: {otros})" if otros else ""))

# ── ADR-003 · un anulado no se reutiliza ────────────────────────────────────
con.execute("UPDATE documento SET estado='anulado', motivo='proveedor equivocado' WHERE id=?", (d1,))
with sdb.transaccion(con): n3 = sdb.reservar_numero(con, "DS")
chk(n3 == ("CTG", 3), "ADR-003: número reutilizado" if n3 != ("CTG", 3) else f"tras anular CTG-1 el siguiente es {n3[0]}-{n3[1]}: el anulado no vuelve")

# ── rango agotado → error claro ─────────────────────────────────────────────
con.execute("UPDATE rango SET ultimo = hasta WHERE clase='DS'")
try:
    with sdb.transaccion(con): sdb.reservar_numero(con, "DS")
    msg = ""
except sdb.RangoAgotado as e: msg = str(e)
chk("CTG" in msg and "10000" in msg and "rango nuevo" in msg, f"rango agotado → error claro: «{msg}»")
try:
    with sdb.transaccion(con): sdb.reservar_numero(con, "NA")
    msg = ""
except sdb.RangoAgotado as e: msg = str(e)
chk("NA" in msg and "no hay rango" in msg, f"clase sin rango (NA, por confirmar) → error claro: «{msg}»")
try: con.execute("INSERT INTO rango(clase, prefijo, desde, hasta, ultimo, activo) VALUES ('DS', 'CTG2', 1, 10, 0, 1)"); dos = True
except sqlite3.IntegrityError: dos = False
chk(not dos, "solo un rango activo por clase")
con.close()

# ── concurrencia: dos hilos con su propia conexión no repiten números ───────
R2 = tmp / "conc.db"; sdb.abrir(R2).close(); vistos, errores = [], []
def reservar(k):
    c = sdb.abrir(R2)
    try:
        for _ in range(k):
            with sdb.transaccion(c): vistos.append(sdb.reservar_numero(c, "DS")[1])
    except Exception as e: errores.append(e)
    finally: c.close()
hs = [threading.Thread(target=reservar, args=(40,)) for _ in range(3)]
for h in hs: h.start()
for h in hs: h.join()
chk(not errores and sorted(vistos) == list(range(1, 121)), f"3 hilos × 40 reservas → 1…120 sin repetidos ni huecos (errores: {errores[:1]})")

# ── T8 · proveedores: registro (R1), alias y encabezados del OCR (R2) ───────
con = sdb.abrir(tmp / "prov.db")
def falla(datos):
    try: prov.guardar(con, datos); return ""
    except prov.DatosIncompletos as e: return str(e)
BASE = dict(tipo_id="CC", num_id="1.047.123.456", nombre="Carlos Cordero", direccion="Bazurto local 12", municipio="Cartagena")
m = falla({**BASE, "tipo_id": "", "num_id": ""})
chk("tipo" in m and "número" in m and not con.execute("SELECT count(*) FROM proveedor").fetchone()[0],
    f"R1: proveedor sin tipo ni número de ID → bloqueado y dice qué falta: «{m}»")
chk("XX" in falla({**BASE, "tipo_id": "XX"}), "R1: tipo de ID que no es de la DIAN → bloqueado")
chk("número" in falla({**BASE, "num_id": "12ab"}), "R1: cédula con letras → bloqueada")
chk("dígito de verificación" in falla({**BASE, "tipo_id": "NIT", "num_id": "900.123.456-0"}), "R1: NIT con DV errado → bloqueado")
chk("dirección" in falla({**BASE, "direccion": " "}) and "correo" in falla({**BASE, "correo": "sin-arroba"}),
    "R1: falta dirección o correo mal escrito → bloqueado")
p1 = prov.guardar(con, BASE)
chk(p1.num_id == "1047123456" and p1.tipo_id == "CC", f"R1: se guarda con el número limpio: {p1.tipo_id} {p1.num_id}")
pn = prov.guardar(con, {**BASE, "tipo_id": "nit", "num_id": "900.123.456-8", "nombre": "Calypso del Caribe S.A."})
chk(pn.tipo_id == "NIT" and pn.num_id == "900123456", f"R1: NIT con DV correcto se guarda sin el DV: {pn.num_id}")
p1b = prov.guardar(con, {**BASE, "num_id": "1047123456", "direccion": "Bazurto local 14"})
chk(p1b.id == p1.id and p1b.direccion == "Bazurto local 14" and con.execute("SELECT count(*) FROM proveedor").fetchone()[0] == 2,
    "R1: guardar el mismo número otra vez actualiza, no duplica")
enc = ["Sin Texto", "[SIN TEXTO]", "Proveedor", "PROVEEDOR:", "Tercero", "Razón social", "Consumidor final", "NIT", "", "  ", "12345", "$ 50.000"]
mal = [t for t in enc if not prov.es_encabezado(t)]
chk(not mal, f"R2: encabezados del OCR reconocidos ({len(enc)})" + (f"; no: {mal}" if mal else ""))
chk(not prov.es_encabezado("Carlos Cordero") and not prov.es_encabezado("Proveedora del Sur"), "R2: un nombre real no es encabezado")
try: prov.asociar(con, "Sin Texto", p1.id); rechaza = False
except prov.DatosIncompletos: rechaza = True
chk(rechaza and prov.resolver(con, "Sin Texto") is None, "R2: «Sin Texto» no se acepta como alias ni se resuelve")
r = prov.resolver(con, "  CARLOS  cordero ")
chk(r is not None and r.id == p1.id, "R2: el nombre guardado ya es alias (mayúsculas y espacios no importan)")
chk(prov.resolver(con, "carlos cordeiro") is None, "R2: alias desconocido → None (la app pide elegir)")
prov.asociar(con, "Cárlos Cordero (Bazurto)", p1.id)
r = prov.resolver(con, "CARLOS CORDERO BAZURTO")
chk(r is not None and r.id == p1.id, "R2: alias asociado se resuelve solo la vez siguiente")
r = prov.resolver(con, "calypso del caribe s.a.")
chk(r is not None and r.id == pn.id, "R2: puntos y tildes no importan en el alias")
try: prov.asociar(con, "carlos cordero bazurto", pn.id); robado = True
except prov.DatosIncompletos: robado = False
chk(not robado and prov.resolver(con, "carlos cordero bazurto").id == p1.id, "R2: un alias de otro proveedor no se cambia sin reasignar=True")
prov.asociar(con, "carlos cordero bazurto", pn.id, reasignar=True)
chk(prov.resolver(con, "carlos cordero bazurto").id == pn.id, "R2: reasignar=True mueve el alias")
con.close()

# ── T9 · semanas: agrupar y plazo ───────────────────────────────────────────
con = sdb.abrir(tmp / "semanas.db")
pc = prov.guardar(con, {"tipo_id": "CC", "num_id": "1047000111", "nombre": "Carlos Cordero", "direccion": "Cra 1 # 2-3", "municipio": "Cartagena"})
def fac(i, f, v, t="Carlos Cordero", **k): return {"id": i, "tipo": "factura", "fecha": f, "valor": v, "tercero": t, "contar": True, "origen": f"p{i}", **k}
def lib(*ds):
    """Un día del libro por fecha de compra (la fecha del papel coincide con la del día)."""
    out = {}
    for d in ds: out.setdefault(d["fecha"], {"fecha": d["fecha"], "documentos": []})["documentos"].append(d)
    return out
L = lib(fac("a", "2026-09-14", 10000), fac("b", "2026-09-16", 25000, "CARLOS  cordero"), fac("c", "2026-09-20", 5000),
        fac("d", "2026-09-21", 7000), fac("z", "2026-09-15", 0), fac("v", "2026-09-15", 900, "Otro", dup_de="a"),
        {"id": "g", "tipo": "gasto_caja", "fecha": "2026-09-15", "valor": 999, "tercero": "Caja"})
gs = sop.semanas(L, con, hoy=date(2026, 9, 10))
g38 = [g for g in gs if g.semana == "2026-W38"]
chk(len(gs) == 2 and len(g38) == 1, f"R3: 3 compras de una semana → 1 grupo (grupos: {[(g.semana, g.total) for g in gs]})")
chk(g38[0].total == 40000 and len(g38[0].compras) == 3 and g38[0].proveedor.id == pc.id, "R3: el total del grupo es 10000+25000+5000 y el alias resuelve al proveedor")
chk(g38[0].lunes == "2026-09-14" and g38[0].domingo == "2026-09-20", "R3: la semana va de lunes a domingo (el domingo 20 cuenta en W38, el 21 en W39)")
chk(sum(len(g.compras) for g in gs) == 4, "R10: duplicados, valor 0 y gastos de caja no entran como compras")
sin = sop.semanas(lib(fac("a", "2026-09-14", 100, "Tienda X"), fac("b", "2026-09-15", 200, "Tienda Y")), con, hoy=date(2026, 9, 10))
chk(len(sin) == 2 and all(g.proveedor is None for g in sin), "R1: terceros sin registrar quedan en grupos aparte, sin proveedor")
enc2 = sop.semanas(lib(fac("a", "2026-09-14", 100, "Sin Texto"), fac("b", "2026-09-15", 200, "SIN TEXTO"), fac("c", "2026-09-15", 300, "")), con, hoy=date(2026, 9, 10))
chk(len(enc2) == 3 and all(g.proveedor is None and g.elegir and len(g.compras) == 1 for g in enc2),
    f"FIX-8/R2: cada compra con tercero ilegible («Sin Texto», vacío) queda en su grupo «elegir proveedor»: {[(g.tercero, g.total, g.elegir) for g in enc2]}")
chk(not any(g.elegir for g in sin), "FIX-8: un tercero legible sin registrar no se marca «elegir»")
chk(len(sop.semanas(L, con, hoy=date(2026, 9, 10), desde="2026-09-21")) == 1, "el filtro desde/hasta recorta las compras")
# FIX-9 · fecha OCR dudosa: a más de `tolerancia_fecha_dias` del día → la del día, con aviso (R6, R10)
mala = sop.semanas({"2026-09-18": {"fecha": "2026-09-18", "documentos": [fac("m", "2018-09-26", 4000), fac("k", "2026-09-16", 1000)]}}, con, hoy=date(2026, 9, 14))
chk(len(mala) == 1 and mala[0].semana == "2026-W38" and [c.fecha for c in mala[0].compras] == ["2026-09-18", "2026-09-16"],
    f"FIX-9: papel 2018-09-26 en el día 2026-09-18 → cae en 2026-W38 con la fecha del día; la de 2 días de diferencia se respeta: {[(g.semana, [c.fecha for c in g.compras]) for g in mala]}")
chk(len(mala[0].avisos) == 1 and "2018-09-26" in mala[0].avisos[0] and "2026-09-18" in mala[0].avisos[0] and mala[0].compras[0].fecha_papel == "2018-09-26",
    f"FIX-9: el grupo avisa con las dos fechas: {mala[0].avisos}")
ocho = {"2026-09-18": {"fecha": "2026-09-18", "documentos": [fac("o", "2026-09-10", 1000)]}}
chk(sop.semanas(ocho, con, hoy=date(2026, 9, 14))[0].compras[0].fecha == "2026-09-18", "FIX-9: 8 días de diferencia > 7 (default) → fecha del día")
chk(sop.semanas(ocho, con, {"tolerancia_fecha_dias": 10}, hoy=date(2026, 9, 14))[0].compras[0].fecha == "2026-09-10"
    and not sop.semanas(ocho, con, {"tolerancia_fecha_dias": 10}, hoy=date(2026, 9, 14))[0].avisos, "FIX-9: tolerancia_fecha_dias configurable (10 → se respeta la del papel, sin aviso)")
try: sop.config({"tolerancia_fecha_dias": -1}); malo = False
except ValueError: malo = True
chk(malo, "FIX-9: tolerancia_fecha_dias negativa → error claro")

# plazos: semana del 14 al 20 de sep 2026; el viernes es el 18
chk(sop.fecha_limite(date(2026, 9, 16), "misma_semana", set()) == date(2026, 9, 18), "R6: sin festivos el límite es el viernes de la semana")
chk(sop.fecha_limite(date(2026, 9, 16), "misma_semana", {date(2026, 9, 18)}) == date(2026, 9, 17), "R6: festivo el viernes → el límite es el jueves")
chk(sop.fecha_limite(date(2026, 9, 14), "misma_semana", {date(2026, 9, 17), date(2026, 9, 18)}) == date(2026, 9, 16), "R6: festivos jueves y viernes → miércoles")
chk(sop.fecha_limite(date(2026, 9, 16), "semana_siguiente", set()) == date(2026, 9, 25), "R6: regla semana_siguiente → viernes de la semana siguiente")
cfg = {"festivos": ["2026-09-18"]}
def de(hoy, c=cfg): return [g for g in sop.semanas(L, con, c, hoy=hoy) if g.semana == "2026-W38"][0]
chk(de(date(2026, 9, 14)).fecha_limite == "2026-09-17", "R6: con el festivo en config el grupo trae el límite 2026-09-17")
chk(de(date(2026, 9, 14)).dias_restantes == 3 and de(date(2026, 9, 14)).estado == sop.EN_PLAZO, "R6: lunes con límite el jueves → quedan 3 hábiles, en plazo")
chk(de(date(2026, 9, 15)).estado == sop.VENCE_PRONTO and de(date(2026, 9, 17)).estado == sop.VENCE_PRONTO, "R6: a 2 días hábiles del límite (y el día mismo) → vence pronto")
chk(de(date(2026, 9, 18)).estado == sop.VENCIDA and de(date(2026, 9, 18)).dias_restantes < 0, "R6: el día después del límite (viernes festivo) → vencida")
chk(de(date(2026, 9, 15), {"festivos": ["2026-09-18"], "aviso_dias": 0}).estado == sop.EN_PLAZO, "R6: aviso_dias configurable (0 → solo avisa el mismo día)")
chk(sop.semanas(L, con, cfg, hoy=date(2026, 9, 30))[0].estado == sop.VENCIDA, "R6: las vencidas salen primero")
try: sop.config({"regla_plazo": "cuando_sea"}); malo = False
except ValueError: malo = True
chk(malo, "regla_plazo desconocida → error claro")
con.close()

# ── T10 · generar: una transacción, consecutivo, sin duplicar (R3, R6, N1, N2) ─
con = sdb.abrir(tmp / "generar.db")
pa = prov.guardar(con, {"tipo_id": "CC", "num_id": "1047000111", "nombre": "Carlos Cordero", "direccion": "Cra 1 # 2-3", "municipio": "Cartagena"})
pb = prov.guardar(con, {"tipo_id": "CC", "num_id": "1047000222", "nombre": "Rosa Pérez", "direccion": "Bazurto 4", "municipio": "Cartagena"})
def dia(f, *ds): return {f: {"fecha": f, "documentos": list(ds)}}
LG = {**dia("2026-09-14", fac("a", "2026-09-14", 10000, numero="F-1"), fac("r", "2026-09-14", 8000, "Rosa Pérez")),
      **dia("2026-09-16", fac("b", "2026-09-16", 25000)), **dia("2026-09-20", fac("c", "2026-09-20", 5000))}
con.execute("UPDATE rango SET ultimo = 57 WHERE clase = 'DS'")
gs = sop.semanas(LG, con, hoy=date(2026, 9, 15))
docs = sop.generar(con, gs, hoy=date(2026, 9, 15))
chk(sorted((d.prefijo, d.numero) for d in docs) == [("CTG", 58), ("CTG", 59)], f"R3: con el último en CTG-57 salen CTG-58 y CTG-59: {[(d.prefijo, d.numero) for d in docs]}")
dc = [d for d in docs if d.proveedor.id == pa.id][0]
chk(len(dc.lineas) == 3 and dc.total == 40000 and dc.semana == "2026-W38", "R3: 3 compras del proveedor en la semana → 1 documento con 3 líneas y total 40000")
fila = con.execute("SELECT total, fecha_generacion, estado, proveedor_id FROM documento WHERE id = ?", (dc.id,)).fetchone()
chk(tuple(fila) == (40000, "2026-09-15", "generado", pa.id), f"el documento queda en la base con fecha de generación y proveedor: {tuple(fila)}")
ls = [tuple(r) for r in con.execute("SELECT fecha_compra, origen_fecha, origen_doc_id, valor, descripcion FROM linea WHERE documento_id = ? ORDER BY fecha_compra", (dc.id,))]
chk([l[:4] for l in ls] == [("2026-09-14", "2026-09-14", "a", 10000), ("2026-09-16", "2026-09-16", "b", 25000), ("2026-09-20", "2026-09-20", "c", 5000)]
    and "F-1" in ls[0][4], f"N2: cada línea guarda fecha, día y id del papel de origen, valor fijo y descripción: {ls[0]}")
LG["2026-09-16"]["documentos"][0]["valor"] = 99999   # editar libro.json después no cambia la copia
chk(con.execute("SELECT valor FROM linea WHERE origen_doc_id = 'b'").fetchone()[0] == 25000, "N2: la línea es una copia fija; editar libro.json no la cambia")
LG["2026-09-16"]["documentos"][0]["valor"] = 25000
n_doc = lambda: con.execute("SELECT count(*) FROM documento").fetchone()[0]
ultimo = lambda: con.execute("SELECT ultimo FROM rango WHERE clase = 'DS'").fetchone()[0]
try: sop.generar(con, gs, hoy=date(2026, 9, 15)); msg = ""
except sop.NoSeGenera as e: msg = str(e)
chk(msg and "CTG-5" in msg and n_doc() == 2 and ultimo() == 59, f"R3: generar 2 veces no duplica ni gasta números: «{msg}»")
otra = sop.semanas(dia("2026-09-21", fac("d", "2026-09-21", 7000)), con, hoy=date(2026, 9, 22))
mezcla = [*otra, *[g for g in sop.semanas(LG, con, hoy=date(2026, 9, 15)) if g.proveedor.id == pb.id]]
try: sop.generar(con, mezcla, hoy=date(2026, 9, 22)); msg = ""
except sop.NoSeGenera as e: msg = str(e)
chk(msg and n_doc() == 2 and ultimo() == 59, "una transacción: si una compra del lote ya está documentada, no se numera ninguna")
con.execute("UPDATE documento SET estado = 'anulado', motivo = 'prueba' WHERE id = ?", (dc.id,))
d2 = sop.generar(con, [g for g in sop.semanas(LG, con, hoy=date(2026, 9, 15)) if g.proveedor.id == pa.id], hoy=date(2026, 9, 15))
chk([(x.prefijo, x.numero) for x in d2] == [("CTG", 60)], f"ADR-003: con el documento anulado sus compras se documentan con un número nuevo: {[(x.prefijo, x.numero) for x in d2]}")
sin = sop.semanas(dia("2026-09-22", fac("t", "2026-09-22", 300, "Tienda Nueva")), con, hoy=date(2026, 9, 22))
try: sop.generar(con, sin, hoy=date(2026, 9, 22)); msg = ""
except sop.NoSeGenera as e: msg = str(e)
chk("Tienda Nueva" in msg and ultimo() == 60, f"R1: un tercero sin registrar no se numera: «{msg}»")
venc = sop.semanas(dia("2026-09-21", fac("d", "2026-09-21", 7000)), con, hoy=date(2026, 10, 5))
try: sop.generar(con, venc, hoy=date(2026, 10, 5)); msg = ""
except sop.NoSeGenera as e: msg = str(e)
chk(venc[0].estado == sop.VENCIDA and "vencid" in msg and ultimo() == 60, f"R6: una semana vencida no se numera sin confirmar: «{msg}»")
dv = sop.generar(con, venc, confirmar_vencidos=True, hoy=date(2026, 10, 5))
chk([(x.numero, x.total) for x in dv] == [(61, 7000)], "R6: con confirmar_vencidos=True la vencida sí se numera")
con.execute("UPDATE rango SET ultimo = hasta WHERE clase = 'DS'")
dos = sop.semanas(dia("2026-09-23", fac("e", "2026-09-23", 100), fac("f", "2026-09-23", 200, "Rosa Pérez")), con, hoy=date(2026, 9, 23))
antes = n_doc()
try: sop.generar(con, dos, hoy=date(2026, 9, 23)); agotado = False
except sdb.RangoAgotado: agotado = True
chk(agotado and n_doc() == antes and not con.execute("SELECT count(*) FROM linea WHERE origen_doc_id IN ('e','f')").fetchone()[0],
    "N1: rango agotado a mitad del lote → error y no queda ningún documento ni línea a medias")
con.close()

# ── FIX-7 · semana ya documentada + compra nueva (R3) ───────────────────────
con = sdb.abrir(tmp / "fix7.db")
prov.guardar(con, {"tipo_id": "CC", "num_id": "1047000111", "nombre": "Carlos Cordero", "direccion": "Cra 1 # 2-3", "municipio": "Cartagena"})
L7 = dia("2026-09-14", fac("a", "2026-09-14", 10000))
k1 = sop.generar(con, sop.semanas(L7, con, hoy=date(2026, 9, 15)), hoy=date(2026, 9, 15))
L7["2026-09-16"] = {"fecha": "2026-09-16", "documentos": [fac("b", "2026-09-16", 4000)]}
g7 = sop.semanas(L7, con, hoy=date(2026, 9, 16))
nuevos, hechos = [g for g in g7 if not g.documentado], [g for g in g7 if g.documentado]
chk(len(nuevos) == 1 and [c.id for c in nuevos[0].compras] == ["b"] and nuevos[0].total == 4000,
    f"R3: la compra nueva de una semana ya documentada forma su propio grupo: {[(g.documentado, [c.id for c in g.compras]) for g in g7]}")
chk(len(hechos) == 1 and hechos[0].documentado == "CTG-1" and [c.id for c in hechos[0].compras] == ["a"], "R3: lo ya documentado sigue saliendo con documentado = CTG-1")
k2 = sop.generar(con, nuevos, hoy=date(2026, 9, 16))
chk([(k1[0].numero, [l.origen_doc_id for l in k1[0].lineas]), (k2[0].numero, [l.origen_doc_id for l in k2[0].lineas])] == [(1, ["a"]), (2, ["b"])],
    f"R3: CTG-1 con a, luego generar da CTG-2 solo con b: {[(d.numero, d.total) for d in k1 + k2]}")
g7 = sop.semanas(L7, con, hoy=date(2026, 9, 16))
chk(sorted(g.documentado for g in g7) == ["CTG-1", "CTG-2"] and all(g.documentado for g in g7), "R3: con las dos documentadas, cada grupo sale con su número")
try: sop.generar(con, g7, hoy=date(2026, 9, 16)); msg = ""
except sop.NoSeGenera as e: msg = str(e)
chk("R3" in msg and con.execute("SELECT ultimo FROM rango").fetchone()[0] == 2, f"R3: un grupo documentado no se numera otra vez: «{msg}»")
con.close()

# varios procesos generan la misma semana a la vez: sale un solo documento
R3 = tmp / "conc_gen.db"; c0 = sdb.abrir(R3)
prov.guardar(c0, {"tipo_id": "CC", "num_id": "1047000111", "nombre": "Carlos Cordero", "direccion": "Cra 1 # 2-3", "municipio": "Cartagena"})
gc = [g for g in sop.semanas(LG, c0, hoy=date(2026, 9, 15)) if g.proveedor]; c0.close(); res = []
def gen():
    c = sdb.abrir(R3)
    try: res.append(len(sop.generar(c, gc, hoy=date(2026, 9, 15))))
    except sop.NoSeGenera: res.append(0)
    finally: c.close()
hs = [threading.Thread(target=gen) for _ in range(4)]
for h in hs: h.start()
for h in hs: h.join()
c0 = sdb.abrir(R3)
chk(sorted(res) == [0, 0, 0, 1] and c0.execute("SELECT count(*) FROM documento").fetchone()[0] == 1,
    f"4 hilos generan la misma semana a la vez → 1 documento, los demás rechazados: {res}")
c0.close()

# ── T11 · soporte_excel.libro ───────────────────────────────────────────────
import io
from openpyxl import load_workbook
from parser_module.contable import soporte_excel as sxl
dv[0].fecha_limite = None   # fecha_limite no se guarda en la base: el RESUMEN debe aguantar None
wb = load_workbook(io.BytesIO(sxl.libro([*d2, *dv])))
chk(wb.sheetnames == ["RESUMEN", "CTG-60", "CTG-61"], f"R5: hoja RESUMEN + una hoja por documento: {wb.sheetnames}")
rs = [tuple(c.value for c in f) for f in wb["RESUMEN"].iter_rows()]
chk(rs[0] == ("Número", "Proveedor", "Semana", "Total", "Fecha límite", "Estado")
    and rs[1][:4] == ("CTG-60", d2[0].proveedor.nombre, "2026-W38", 40000) and rs[1][5] == "generado"
    and rs[2][0] == "CTG-61" and rs[2][3] == 7000 and rs[2][4] is None, f"RESUMEN: número, proveedor, semana, total, fecha límite, estado: {rs[1:]}")
h = wb["CTG-60"]; campos = {f[0].value: f[1].value for f in h.iter_rows() if f[0].value}
chk(campos.get("Número") == "CTG-60" and campos.get("Fecha de generación") == "2026-09-15" and campos.get("Tipo de documento del vendedor") == "CC"
    and campos.get("Número de documento") == d2[0].proveedor.num_id and campos.get("Nombre del vendedor") == d2[0].proveedor.nombre
    and campos.get("Dirección") and campos.get("Municipio"), "art. 4: número, fecha de generación y datos del vendedor en la hoja del documento")
filas = [tuple(c.value for c in f) for f in h.iter_rows()]
det = [f for f in filas if f[0] in ("2026-09-14", "2026-09-16", "2026-09-20")]
chk(len(det) == 3 and sum(f[2] for f in det) == 40000 and filas[-2][1:3] == ("IVA", 0) and filas[-1][1:3] == ("Total", 40000),
    f"art. 4: detalle de la operación (3 líneas con papel de origen), IVA y total: {filas[-2:]}")


# ── T12 · rutas /api/soporte/* y /api/proveedores ───────────────────────────
import json, app as web
T12 = Path(tempfile.mkdtemp())
web.RUTA_SOPORTE, web.RUTA_LIBRO, web.RUTA_CONT = T12 / "soporte.db", T12 / "libro.json", T12 / "contable.json"
web.RUTA_LIBRO.write_text(json.dumps({"2026-09-14": {"fecha": "2026-09-14", "documentos": [
    {"id": "f1", "tipo": "factura", "fecha": "2026-09-14", "valor": 12000, "tercero": "Carlos Cordero", "origen": "foto1.jpg"},
    {"id": "f2", "tipo": "factura", "fecha": "2026-09-15", "valor": 8000, "tercero": "Carlos Cordero", "origen": "foto2.jpg"}]}}), encoding="utf8")
cl = web.app.test_client()
AJENO = {"Origin": "http://evil.example.com"}
r = cl.post("/api/proveedores", json=BASE, headers=AJENO)
chk(r.status_code == 403, f"T1: POST con Origin ajeno a /api/proveedores → 403 ({r.status_code})")
r = cl.post("/api/soporte/generar", json={"semanas": []}, headers=AJENO)
chk(r.status_code == 403, f"T1: POST con Origin ajeno a /api/soporte/generar → 403 ({r.status_code})")
r = cl.post("/api/proveedores", json=BASE, headers={"Host": "evil.example.com"})
chk(r.status_code == 403, f"T1: POST con Host ajeno (DNS rebinding) → 403 ({r.status_code})")
for u in ("/api/proveedores", "/api/soporte/excel", "/api/soporte/semanas"):
    r = cl.get(u, headers={"Host": "evil.example.com"})
    chk(r.status_code == 403, f"FIX-10: GET {u} con Host ajeno (DNS rebinding) → 403 ({r.status_code})")
r = cl.get("/api/proveedores", headers={"Host": "[::1]:5000"})
chk(r.status_code == 200, f"FIX-10: GET con Host [::1] (IPv6 local) → 200 ({r.status_code})")
r = cl.get("/api/proveedores", headers={"Host": "127.0.0.1:5000"})
chk(r.status_code == 200, f"FIX-10: GET con Host 127.0.0.1 → 200 ({r.status_code})")
r = cl.post("/api/soporte/generar", json={"semanas": []}, headers={"Host": "[::1]:5000", "Origin": "http://[::1]:5000"})
chk(r.status_code != 403, f"FIX-10: POST desde [::1] con su propio Origin no da 403 ({r.status_code})")
c0 = sdb.abrir(web.RUTA_SOPORTE)
chk(c0.execute("SELECT count(*) FROM proveedor").fetchone()[0] == 0 and c0.execute("SELECT ultimo FROM rango").fetchone()[0] == 0,
    "T1: los POST rechazados no escribieron nada"); c0.close()
LOC = {"Origin": "http://localhost"}
r = cl.post("/api/proveedores", json={**BASE, "direccion": ""}, headers=LOC)
chk(r.status_code == 400 and "dirección" in r.get_json()["error"], f"R1: proveedor sin dirección → 400 con el motivo: {r.get_json()}")
r = cl.post("/api/proveedores", json={**BASE, "alias_ocr": "Carlos Cordero"}, headers=LOC)
pid = r.get_json()["proveedor"]["id"] if r.status_code == 200 else None
chk(r.status_code == 200 and pid, f"flujo feliz: guardar proveedor → 200 (id {pid})")
chk([p["nombre"] for p in cl.get("/api/proveedores").get_json()["proveedores"]] == ["Carlos Cordero"], "GET /api/proveedores lista el guardado")
r = cl.get("/api/soporte/semanas")
sm = r.get_json()["semanas"] if r.status_code == 200 else []
chk(len(sm) == 1 and sm[0]["total"] == 20000 and sm[0]["registrado"] and sm[0]["proveedor_id"] == pid and len(sm[0]["compras"]) == 2,
    f"GET /api/soporte/semanas: 1 grupo de 20000 con 2 compras: {[(x['semana'], x['total']) for x in sm]}")
chk(cl.get("/api/soporte/semanas?desde=mal").status_code == 400, "semanas: fecha mal escrita → 400")
sel = {"semana": sm[0]["semana"], "proveedor_id": pid}
r = cl.post("/api/soporte/generar", json={"semanas": [sel]}, headers=LOC)
chk(r.status_code == 409 and "R6" in r.get_json()["error"], f"R6: semana vencida sin confirmar → 409: {r.get_json()}")
r = cl.post("/api/soporte/generar", json={"semanas": [{"semana": "2020-W01", "proveedor_id": pid}]}, headers=LOC)
chk(r.status_code == 400, "semana que no existe → 400")
r = cl.post("/api/soporte/generar", json={"semanas": [sel], "confirmar_vencidos": True}, headers=LOC)
dj = r.get_json()["documentos"] if r.status_code == 200 else []
chk(r.status_code == 200 and len(dj) == 1 and dj[0]["numero"] == "CTG-1" and dj[0]["total"] == 20000, f"flujo feliz: generar → 200, CTG-1 por 20000: {dj and dj[0]['numero']}")
r = cl.post("/api/soporte/generar", json={"semanas": [sel], "confirmar_vencidos": True}, headers=LOC)
c0 = sdb.abrir(web.RUTA_SOPORTE)
chk(r.status_code == 409 and c0.execute("SELECT count(*) FROM documento").fetchone()[0] == 1 and c0.execute("SELECT ultimo FROM rango").fetchone()[0] == 1,
    f"R3: generar otra vez no duplica ni gasta número ({r.status_code}): {r.get_json()}"); c0.close()
chk(sm[0]["documentado"] is None, "semanas: antes de generar, documentado es null")
sm2 = cl.get("/api/soporte/semanas").get_json()["semanas"]
chk(len(sm2) == 1 and sm2[0]["documentado"] == "CTG-1", f"semanas: tras generar, la semana sale marcada con CTG-1: {[x['documentado'] for x in sm2]}")
r = cl.get("/api/soporte/excel")
wb =load_workbook(io.BytesIO(r.data)) if r.status_code == 200 else None
rs = [tuple(c.value for c in f) for f in wb["RESUMEN"].iter_rows()] if wb else []
chk(r.status_code == 200 and wb.sheetnames == ["RESUMEN", "CTG-1"] and rs[1][0] == "CTG-1" and rs[1][4] and rs[1][3] == 20000,
    f"GET /api/soporte/excel → 200, hojas RESUMEN y CTG-1, con fecha límite recalculada: {rs[1:]}")
chk(cl.get("/api/soporte/excel?ids=999").status_code == 404 and cl.get("/api/soporte/excel?ids=x").status_code == 400, "excel: ids sin documentos → 404, mal escritos → 400")
lj = json.loads(web.RUTA_LIBRO.read_text(encoding="utf8"))
lj["2026-09-14"]["documentos"].append({"id": "f3", "tipo": "factura", "fecha": "2026-09-16", "valor": 3000, "tercero": "Carlos Cordero", "origen": "foto3.jpg"})
web.RUTA_LIBRO.write_text(json.dumps(lj), encoding="utf8")
sm3 = cl.get("/api/soporte/semanas").get_json()["semanas"]
chk(sorted((x["documentado"] or "", x["total"]) for x in sm3) == [("", 3000), ("CTG-1", 20000)], f"FIX-7: la compra nueva sale aparte de la semana CTG-1: {[(x['documentado'], x['total']) for x in sm3]}")
r = cl.post("/api/soporte/generar", json={"semanas": [sel], "confirmar_vencidos": True}, headers=LOC)
dj = r.get_json().get("documentos", []) if r.status_code == 200 else []
chk(r.status_code == 200 and [(d["numero"], d["total"]) for d in dj] == [("CTG-2", 3000)], f"FIX-7: generar la misma semana da CTG-2 solo con la nueva: {r.get_json()}")

# ── T14 · compra manual (R4, N2): queda en libro.json como factura «manual», sale en DATOS OCR y en las semanas ──
from parser_module.contable import construir as construir_asientos
CM = {"tercero": "Carlos Cordero", "fecha": "2026-09-22", "descripcion": "Verduras pagadas por transferencia", "valor": 15000, "medio": "Transferencia"}
r = cl.post("/api/soporte/compra", json=CM, headers=AJENO)
chk(r.status_code == 403, f"T1: POST compra manual con Origin ajeno → 403 ({r.status_code})")
for malo, motivo in (({"valor": 0}, "valor 0"), ({"valor": "x"}, "valor no numérico"), ({"valor": 10**12}, "valor absurdo"), ({"fecha": "2026-02-30"}, "fecha inexistente"),
                     ({"tercero": "  "}, "sin proveedor"), ({"tercero": "Sin Texto"}, "proveedor rótulo"), ({"descripcion": ""}, "sin descripción"),
                     ({"descripcion": "x" * 300}, "descripción larga"), ({"medio": "Bitcoin"}, "medio desconocido")):
    r = cl.post("/api/soporte/compra", json={**CM, **malo}, headers=LOC)
    chk(r.status_code == 400 and r.get_json().get("error"), f"R4: compra manual con {motivo} → 400: {r.get_json()}")
chk("2026-09-22" not in json.loads(web.RUTA_LIBRO.read_text(encoding="utf8")), "las compras rechazadas no escribieron en libro.json")
r = cl.post("/api/soporte/compra", json=CM, headers=LOC)
chk(r.status_code == 200, f"flujo feliz: compra manual → 200 ({r.get_json()})")
lj = json.loads(web.RUTA_LIBRO.read_text(encoding="utf8"))
dm = [d for d in lj.get("2026-09-22", {}).get("documentos", []) if d.get("origen") == "manual"]
chk(len(dm) == 1 and dm[0]["tipo"] == "factura" and dm[0]["valor"] == 15000 and dm[0]["fecha"] == "2026-09-22" and dm[0]["tercero"] == "Carlos Cordero",
    f"libro.json: el día 2026-09-22 guarda la factura manual: {dm and {k: dm[0][k] for k in ('tipo', 'origen', 'valor', 'fecha', 'tercero')}}")
sm4 = [x for x in cl.get("/api/soporte/semanas").get_json()["semanas"] if x["semana"] == "2026-W39"]
chk(len(sm4) == 1 and sm4[0]["total"] == 15000 and sm4[0]["registrado"] and [c["origen"] for c in sm4[0]["compras"]] == ["manual"]
    and sm4[0]["compras"][0]["descripcion"] == CM["descripcion"], f"R4: la compra manual queda en el acumulado de su semana (W39) con origen «manual»: {sm4}")
dia22 = lj["2026-09-22"]
rd = construir_asientos("2026-09-22", dia22["documentos"], dia22.get("banco"), {}, dia22.get("cierre"))
fm = [x for x in rd["datos"] if x["archivo"] == "manual"]
chk(len(fm) == 1 and fm[0]["estado"] == "ok" and fm[0]["valor"] == 15000 and fm[0]["tercero"] == "Carlos Cordero" and fm[0]["medio"] == "Transferencia",
    f"R9: la compra manual sale en DATOS OCR con archivo «manual»: {fm}")
chk((dia22["debito"], dia22["credito"]) == (rd["debito"], rd["credito"]) and rd["debito"] > 0, f"el día guardado queda recalculado (sin aviso de FIX-6): {dia22['debito']} / {rd['debito']}")
from parser_module.contable.documentos import _base as base_doc
lj["2026-09-14"]["documentos"] = [{**base_doc("factura", ""), **d} for d in lj["2026-09-14"]["documentos"]]   # papeles leídos traen todos los campos de _base
web.RUTA_LIBRO.write_text(json.dumps(lj), encoding="utf8")
r = cl.post("/api/soporte/compra", json={**CM, "fecha": "2026-09-14", "valor": 4000}, headers=LOC)
lj = json.loads(web.RUTA_LIBRO.read_text(encoding="utf8"))
ids14 = [d["id"] for d in lj["2026-09-14"]["documentos"]]
chk(r.status_code == 200 and ids14[:3] == ["f1", "f2", "f3"] and len(ids14) == 4, f"en un día que ya existe la compra se agrega sin tocar los papeles leídos: {ids14}")

# ── T15 · ciclo de vida: transmitido, anulado, bloqueos (R7, N1) y datos del vendedor fijos ──
c15 = sdb.abrir(tmp / "t15.db")
pc = prov.guardar(c15, {"tipo_id": "CC", "num_id": "1047000333", "nombre": "Carlos Cordero", "direccion": "Cra 1 # 2-3", "municipio": "Cartagena"})
prov.guardar(c15, {"tipo_id": "CC", "num_id": "1047000444", "nombre": "Rosa Pérez", "direccion": "Bazurto 4", "municipio": "Cartagena"})
L15 = {**dia("2026-09-14", fac("a", "2026-09-14", 10000), fac("r", "2026-09-14", 8000, "Rosa Pérez")), **dia("2026-09-21", fac("n", "2026-09-21", 3000))}
H15 = date(2026, 9, 15)
g15 = [g for g in sop.semanas(L15, c15, hoy=H15) if g.semana == "2026-W38"]
d15 = sop.generar(c15, g15, hoy=H15)
dc = next(d for d in d15 if d.proveedor.nombre == "Carlos Cordero"); dr = next(d for d in d15 if d.proveedor.nombre == "Rosa Pérez")
prov.guardar(c15, {"tipo_id": "CC", "num_id": "1047000333", "nombre": "Carlos A. Cordero", "direccion": "Calle Nueva 9", "municipio": "Turbaco"})
lc = next(d for d in sop.documentos(c15) if d.id == dc.id)
chk((lc.proveedor.nombre, lc.proveedor.direccion, lc.proveedor.municipio) == ("Carlos Cordero", "Cra 1 # 2-3", "Cartagena") and lc.proveedor.id == pc.id,
    f"art. 4: el documento conserva los datos del vendedor del día de generación aunque el proveedor se edite después: {lc.proveedor}")
h15 = load_workbook(io.BytesIO(sxl.libro(sop.documentos(c15, ids=[dc.id]))))[sxl.nombre_hoja(dc)]
cp = {f[0].value: f[1].value for f in h15.iter_rows() if f[0].value}
chk(cp.get("Nombre del vendedor") == "Carlos Cordero" and cp.get("Municipio") == "Cartagena", f"el Excel sale con los datos fijos del vendedor: {cp.get('Nombre del vendedor')}, {cp.get('Municipio')}")
# migración 2: documentos de antes de la copia se llenan con los datos de ese momento
viejo = sdb.MIGRACIONES; sdb.MIGRACIONES = viejo[:1]
cm = sdb.abrir(tmp / "mig.db"); sdb.MIGRACIONES = viejo
pm = prov.guardar(cm, {"tipo_id": "CC", "num_id": "1047000555", "nombre": "Ana Díaz", "direccion": "Getsemaní 1", "municipio": "Cartagena"})
cm.execute("INSERT INTO documento(clase, prefijo, numero, proveedor_id, semana_iso, fecha_generacion, total, creado) VALUES ('DS','CTG',1,?,'2026-W38','2026-09-15',100,'x')", (pm.id,))
sdb.migrar(cm)
dm = sop.documentos(cm)
chk(cm.execute("PRAGMA user_version").fetchone()[0] == len(sdb.MIGRACIONES) and dm[0].proveedor.nombre == "Ana Díaz" and dm[0].proveedor.no_obligado is True,
    f"migración 2 (expand + llenado): un documento viejo queda con la copia del vendedor: {dm[0].proveedor}"); cm.close()

def no_se(f, *a):
    try: f(*a); return ""
    except sop.NoSePuede as e: return str(e)
chk(no_se(sop.marcar_transmitido, c15, dc.id, "abc") and sop.documentos(c15, ids=[dc.id])[0].estado == "generado", "R7: un CUDS mal formado se rechaza y no cambia el estado")
CUDS = "AB" * 48
chk(no_se(sop.marcar_transmitido, c15, dc.id, CUDS) == "", "R7: marcar transmitido con CUDS de 96 hexadecimales")
lt = sop.documentos(c15, ids=[dc.id])[0]
chk((lt.estado, lt.cuds) == ("transmitido", CUDS.lower()), f"queda transmitido con su CUDS: {lt.estado} {lt.cuds[:8]}…")
chk("generado" in no_se(sop.marcar_transmitido, c15, dc.id), "R7: transmitido dos veces → error")
chk("ADR-003: transmitido no se anula" in no_se(sop.anular, c15, dc.id, "error"), "ADR-003: anular un transmitido → error (se corrige con nota de ajuste)")
chk(no_se(sop.marcar_transmitido, c15, 9999) != "", "R7: documento inexistente → error")
chk(no_se(sop.marcar_transmitido, c15, dr.id, None) == "" and sop.documentos(c15, ids=[dr.id])[0].estado == "transmitido"
    and sop.documentos(c15, ids=[dr.id])[0].cuds == "", "R7: el CUDS es opcional")

# bloqueos: lo transmitido no cambia desde el libro
F14 = L15["2026-09-14"]["documentos"]
chk(sop.bloqueos(c15, "2026-09-14", F14) == [], "R7: guardar el día sin tocar las compras transmitidas → sin bloqueos")
chk(sop.bloqueos(c15, "2026-09-14", [{**F14[0], "descripcion": "otra", "cuenta": "51"}, F14[1]]) == [], "R7: cambiar otros campos (cuenta, descripción) no bloquea")
b = sop.bloqueos(c15, "2026-09-14", [{**F14[0], "valor": 12000}, F14[1]])
chk(len(b) == 1 and sxl.nombre_hoja(dc) in b[0] and "10.000" in b[0] and "nota de ajuste" in b[0], f"R7: editar el valor de una compra transmitida → bloqueado: {b}")
b = sop.bloqueos(c15, "2026-09-14", [{**F14[0], "tercero": "Rosa Pérez"}, F14[1]])
chk(len(b) == 1 and "proveedor" in b[0], f"R7: cambiar el tercero a otro proveedor → bloqueado: {b}")
chk(sop.bloqueos(c15, "2026-09-14", [{**F14[0], "tercero": "  CARLOS CORDERO "}, F14[1]]) == [], "el mismo proveedor escrito distinto (alias normalizado) no bloquea")
chk(len(sop.bloqueos(c15, "2026-09-14", [F14[1]])) == 1, "R7: quitar la compra del día → bloqueado")
chk(len(sop.bloqueos(c15, "2026-09-14", [{**F14[0], "contar": False}, F14[1]])) == 1, "R7: dejar de contar la compra → bloqueado")
chk(len(sop.bloqueos(c15, "2026-09-14", [])) == 2, "R7: borrar el día → bloquea las 2 compras transmitidas")

# anular: el número queda gastado, sale en el RESUMEN y no vuelve (N1, ADR-003)
g21 = [g for g in sop.semanas(L15, c15, hoy=date(2026, 9, 22)) if g.semana == "2026-W39"]
dn = sop.generar(c15, g21, hoy=date(2026, 9, 22))[0]
chk(sop.bloqueos(c15, "2026-09-21", [{**L15["2026-09-21"]["documentos"][0], "valor": 1}]) == [], "un documento solo generado (no transmitido) no bloquea el libro")
chk("motivo" in no_se(sop.anular, c15, dn.id, "  "), "anular sin motivo → error")
chk(no_se(sop.anular, c15, dn.id, "Valor mal leído") == "", "anular un generado con motivo")
chk("anulado" in no_se(sop.anular, c15, dn.id, "otra vez") and "anulado" in no_se(sop.marcar_transmitido, c15, dn.id), "un anulado no se anula otra vez ni se transmite")
g21b = [g for g in sop.semanas(L15, c15, hoy=date(2026, 9, 22)) if g.semana == "2026-W39" and not g.documentado]
chk(len(g21b) == 1 and g21b[0].compras[0].id == "n", "las compras de un anulado quedan libres para un documento nuevo")
dn2 = sop.generar(c15, g21b, hoy=date(2026, 9, 22))[0]
nums = [d.numero for d in sop.documentos(c15)]
chk(dn2.numero == dn.numero + 1 and len(set(nums)) == len(nums), f"N1: el documento nuevo toma el siguiente número, el anulado no vuelve: {dn.numero} anulado → {dn2.numero}; todos {nums}")
wb15 = load_workbook(io.BytesIO(sxl.libro(sop.documentos(c15))))
rs15 = {f[0].value: f[5].value for f in wb15["RESUMEN"].iter_rows(min_row=2)}
chk(rs15.get(f"CTG-{dn.numero}") == "anulado" and rs15.get(f"CTG-{dn2.numero}") == "generado" and rs15.get(sxl.nombre_hoja(dc)) == "transmitido",
    f"N1: el anulado aparece en el RESUMEN como «anulado»: {rs15}")
chk(sop.documentos(c15, ids=[dn.id])[0].motivo == "Valor mal leído", "el motivo de la anulación queda guardado")

# rutas
T15 = Path(tempfile.mkdtemp())
web.RUTA_SOPORTE, web.RUTA_LIBRO, web.RUTA_CONT = T15 / "soporte.db", T15 / "libro.json", T15 / "contable.json"
FA = {**base_doc("factura", ""), "id": "fa", "fecha": "2026-09-14", "valor": 12000, "tercero": "Carlos Cordero", "origen": "foto1.jpg", "medio": "Efectivo"}
r = cl.post("/api/contable/libro", json={"fecha": "2026-09-14", "documentos": [FA]}, headers=LOC)
chk(r.status_code == 200 and not web.RUTA_SOPORTE.exists(), f"sin soporte.db, guardar el día funciona y no crea la base ({r.status_code})")
r = cl.post("/api/proveedores", json=BASE, headers=LOC)
r = cl.post("/api/soporte/generar", json={"semanas": [{"semana": "2026-W38", "proveedor_id": r.get_json()["proveedor"]["id"]}], "confirmar_vencidos": True}, headers=LOC)
did = r.get_json()["documentos"][0]["id"]
r = cl.post(f"/api/soporte/{did}/transmitido", json={"cuds": "x"}, headers=LOC)
chk(r.status_code == 409 and "CUDS" in r.get_json()["error"], f"POST transmitido con CUDS malo → 409: {r.get_json()}")
r = cl.post(f"/api/soporte/{did}/transmitido", json={"cuds": CUDS}, headers=AJENO)
chk(r.status_code == 403, f"T1: POST transmitido con Origin ajeno → 403 ({r.status_code})")
r = cl.post(f"/api/soporte/{did}/transmitido", json={"cuds": CUDS}, headers=LOC)
chk(r.status_code == 200, f"POST transmitido → 200 ({r.get_json()})")
dj = cl.get("/api/soporte/documentos").get_json()["documentos"]
chk(len(dj) == 1 and dj[0]["estado"] == "transmitido" and dj[0]["cuds"] == CUDS.lower() and dj[0]["numero"] == "CTG-1", f"GET documentos lista estado y CUDS: {dj and {k: dj[0][k] for k in ('numero', 'estado')}}")
antes = web.RUTA_LIBRO.read_text(encoding="utf8")
r = cl.post("/api/contable/libro", json={"fecha": "2026-09-14", "documentos": [{**FA, "valor": 15000}], "reemplazar": True}, headers=LOC)
chk(r.status_code == 409 and len(r.get_json().get("bloqueos", [])) == 1 and web.RUTA_LIBRO.read_text(encoding="utf8") == antes,
    f"R7: reemplazar el día con el valor de una compra transmitida cambiado → 409 y libro.json intacto: {r.get_json()}")
r = cl.delete("/api/contable/libro/2026-09-14", headers=LOC)
chk(r.status_code == 409 and "2026-09-14" in json.loads(web.RUTA_LIBRO.read_text(encoding="utf8")), f"R7: borrar un día con compras transmitidas → 409 ({r.status_code})")
r = cl.post("/api/contable/libro", json={"fecha": "2026-09-14", "documentos": [{**FA, "cuenta": "51359501"}], "reemplazar": True}, headers=LOC)
chk(r.status_code == 200, f"reemplazar el día sin tocar valor ni tercero de lo transmitido → 200 ({r.get_json()})")
r = cl.post(f"/api/soporte/{did}/anular", json={"motivo": "x"}, headers=LOC)
chk(r.status_code == 409 and "ADR-003" in r.get_json()["error"], f"POST anular un transmitido → 409: {r.get_json()}")
r = cl.post("/api/soporte/9999/anular", json={"motivo": "x"}, headers=LOC)
chk(r.status_code == 409, f"POST anular inexistente → 409 ({r.status_code})")
r = cl.post(f"/api/soporte/{did}/anular", json={}, headers=LOC)
chk(r.status_code == 400, f"POST anular sin motivo → 400 ({r.status_code})")

print("\n" + ("TODO OK" if ok else "HAY FALLOS"))
