"""Genera el Excel final: hoja «Ingresos» (una fila por día; columnas calculadas y controles como FÓRMULAS vivas)
y «Resumen mensual» (sumas por mes, y ganancias/pérdidas por separado)."""
from __future__ import annotations
import io
from datetime import date
from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from parser_module import Esquema
from parser_module.formulas import a_excel

AZUL, AZUL_OSC, AZUL_CLARO = "D9ECFC", "17466F", "EEF6FE"
_fill = lambda c: PatternFill("solid", fgColor=c)
_borde = Border(bottom=Side(style="thin", color="CFE3F6"))
FMT = {"monto": '#,##0;-#,##0', "conteo": '0', "calculo": '#,##0;-#,##0'}
FMT_RES = '[Color10]+#,##0;[Red]-#,##0;"cuadrado"'


def _fecha(s):
    try: return date.fromisoformat(s)
    except (TypeError, ValueError): return s or ""


def _cab(ws, fila=1):
    for c in ws[fila]:
        c.font = Font(bold=True, color=AZUL_OSC); c.fill = _fill(AZUL)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[fila].height = 34


def construir_excel(filas: list[dict], esquema: Esquema | None = None) -> bytes:
    esq = esquema or Esquema.defecto()
    cols = esq.columnas
    filas = sorted(filas, key=lambda f: f.get("fecha") or "9999")
    wb = Workbook(); ws = wb.active; ws.title = "Ingresos"
    col_de = {c.clave: i for i, c in enumerate(cols, 2)}            # clave → número de columna
    c_ctrl = len(cols) + 2 if esq.controles else None
    c_nota = len(cols) + 2 + (1 if esq.controles else 0)
    ws.append(["Fecha"] + [c.nombre for c in cols] + (["Controles"] if c_ctrl else []) + ["Notas"])
    _cab(ws)
    n = len(filas)
    for r, f in enumerate(filas, 2):
        v = f.get("valores") or {}
        fila = [_fecha(f.get("fecha"))]
        for c in cols:
            if c.tipo == "calculo":
                fila.append("=IFERROR(" + a_excel(c.formula, lambda k: f"{L(col_de[k])}{r}") + ",0)")
            else:
                fila.append(int(v.get(c.clave) or 0))
        if c_ctrl:
            conds = [f"ABS(IFERROR({a_excel(x['formula'], lambda k: f'{L(col_de[k])}{r}')},0))<={x['tolerancia']}" for x in esq.controles]
            fila.append(f'=IF(AND({",".join(conds)}),"✓","⚠ revisar")')
        fila.append(f.get("notas", ""))
        ws.append(fila)
        ws.cell(r, 1).number_format = "dd/mm/yyyy"; ws.cell(r, 1).alignment = Alignment(horizontal="center")
        for c in cols:
            cell = ws.cell(r, col_de[c.clave])
            cell.number_format = FMT_RES if (c.tipo == "calculo" and c.resultado) else FMT[c.tipo]
            if c.tipo == "calculo": cell.font = Font(bold=True)
        if c_ctrl: ws.cell(r, c_ctrl).alignment = Alignment(horizontal="center")
        for cell in ws[r]: cell.border = _borde
    t = n + 2
    ws.cell(t, 1, "TOTAL")
    for c in cols:
        i = col_de[c.clave]
        cell = ws.cell(t, i, f"=SUM({L(i)}2:{L(i)}{t - 1})" if n else 0)
        cell.number_format = FMT_RES if (c.tipo == "calculo" and c.resultado) else FMT[c.tipo]
    for cell in ws[t]: cell.font = Font(bold=True, color=AZUL_OSC); cell.fill = _fill(AZUL_CLARO)
    ws.column_dimensions["A"].width = 13
    for i in range(2, len(cols) + 2): ws.column_dimensions[L(i)].width = 17
    if c_ctrl:
        ws.column_dimensions[L(c_ctrl)].width = 13
        if n: ws.conditional_formatting.add(f"{L(c_ctrl)}2:{L(c_ctrl)}{n + 1}",
                                            FormulaRule(formula=[f'LEFT({L(c_ctrl)}2,1)="⚠"'], fill=_fill("FFE08A")))
    ws.column_dimensions[L(c_nota)].width = 48
    ws.freeze_panes = "B2"
    if n: ws.auto_filter.ref = f"A1:{L(c_nota)}{n + 1}"

    # ── Resumen mensual (SUMIFS sobre «Ingresos»; para ganancia/pérdida, también por separado) ──
    rs = wb.create_sheet("Resumen mensual")
    spec = [("Días", None, "cnt")]
    for c in cols:
        spec.append((c.nombre, c, "sum"))
        if c.tipo == "calculo" and c.resultado:
            spec += [(f"{c.nombre} · ganancias", c, "pos"), (f"{c.nombre} · pérdidas", c, "neg")]
    rs.append(["Mes"] + [s[0] for s in spec]); _cab(rs)
    meses = sorted({f["fecha"][:7] for f in filas if isinstance(_fecha(f.get("fecha")), date)})
    for j, m in enumerate(meses, 2):
        rs.cell(j, 1, date(int(m[:4]), int(m[5:7]), 1)).number_format = "mmmm yyyy"
        crit = f'Ingresos!$A:$A,">="&$A{j},Ingresos!$A:$A,"<"&EDATE($A{j},1)'
        for i, (nom, c, tipo) in enumerate(spec, 2):
            if tipo == "cnt": cell = rs.cell(j, i, f"=COUNTIFS({crit})"); cell.number_format = "0"; continue
            col = L(col_de[c.clave]); rango = f"Ingresos!${col}:${col}"
            extra = {"sum": "", "pos": f',{rango},">0"', "neg": f',{rango},"<0"'}[tipo]
            cell = rs.cell(j, i, f"=SUMIFS({rango},{crit}{extra})")
            cell.number_format = '#,##0;[Red]-#,##0' if tipo != "sum" or c.tipo == "calculo" else FMT[c.tipo]
        for cell in rs[j]: cell.border = _borde
    if meses:
        tr = len(meses) + 2; rs.cell(tr, 1, "TOTAL")
        for i in range(2, len(spec) + 2):
            cell = rs.cell(tr, i, f"=SUM({L(i)}2:{L(i)}{tr - 1})"); cell.number_format = '#,##0;[Red]-#,##0'
        for cell in rs[tr]: cell.font = Font(bold=True, color=AZUL_OSC); cell.fill = _fill(AZUL_CLARO)
    rs.column_dimensions["A"].width = 18
    for i in range(2, len(spec) + 2): rs.column_dimensions[L(i)].width = 17
    rs.freeze_panes = "B2"
    # ── Detalle por factura: de qué imágenes salió cada día (para auditar) ──
    det = [(f, x) for f in filas for x in (f.get("facturas") or [])]
    if det:
        ds = wb.create_sheet("Detalle por factura")
        ent = [c.clave for c in esquema.entradas]
        ds.append(["Fecha", "Factura"] + [esquema.nombres[k] for k in ent] + ["Notas"])
        _cab(ds)
        for f, x in det:
            xv = x.get("valores") or {}
            ds.append([_fecha(f.get("fecha")), x.get("nombre", "")] + [int(xv.get(k) or 0) for k in ent] + [x.get("notas", "")])
        for r in range(2, len(det) + 2):
            ds.cell(r, 1).number_format = "dd/mm/yyyy"
            for i, k in enumerate(ent, 3): ds.cell(r, i).number_format = '#,##0;[Red]-#,##0'
        ds.column_dimensions["A"].width = 13; ds.column_dimensions["B"].width = 30
        for i in range(3, len(ent) + 3): ds.column_dimensions[L(i)].width = 16
        ds.column_dimensions[L(len(ent) + 3)].width = 40
        ds.freeze_panes = "C2"
    buf = io.BytesIO(); wb.save(buf)
    return buf.getvalue()
