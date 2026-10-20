"""Prueba del libro contable con el texto que leería el OCR de TUS papeles (recibo de venta, gastos de caja, factura)."""
import io, sys
from parser_module.contable import abreviar_recibos, construir, leer_texto, rangos, libro_a_excel, Clasificador

from parser_module.contable.ejemplos import FOTO_PAPELES, GASTO, VENTA
ok = True
def chk(c, m):
    global ok; ok &= bool(c); print(("✓ " if c else "✗ ") + m)

docs = leer_texto(FOTO_PAPELES) + leer_texto(VENTA)
chk([d["tipo"] for d in docs] == ["gasto_caja", "caja_menor_manuscrito", "gasto_caja", "caja_menor_manuscrito", "gasto_caja", "caja_menor_manuscrito", "factura", "venta"],
    f"8 papeles reconocidos (6 de una sola foto): {[d['tipo'][:6] for d in docs]}")
v = docs[-1]
chk((v["numero"], v["consumo"], v["propina"], v["efectivo"], v["tarjeta"], v["fecha"]) == ("85531", 178000, 17800, 100000, 95800, "2026-09-12"), f"recibo de venta: {v['numero']} · consumo {v['consumo']} · propina {v['propina']} · efectivo {v['efectivo']} · tarjeta {v['tarjeta']}")
chk(not v["avisos"], "los pagos cuadran con consumo + propina")
g = {d["numero"]: d for d in docs if d["tipo"] == "gasto_caja"}
chk((g["220"]["valor"], g["217"]["valor"], g["216"]["valor"]) == (26000, 86000, 2800) and "transportes de compras" in g["220"]["detalle"], "gastos de caja: valor, consecutivo y descripción")
f = docs[6]
chk((f["valor"], f["numero"], f["tercero"], f["medio"]) == (10000, "POS67085", "Sebastian Hernandez De La Espriella", "Transferencia"), f"factura: {f['tercero']} · {f['valor']} · {f['medio']} (la nota a mano «Transf» manda sobre «Efectivo»)")

cl = Clasificador()
for d in docs:
    if d["tipo"] in ("gasto_caja", "factura"):
        r = cl.clasificar(d["tercero"] if d["tipo"] == "factura" else "", d["detalle"], d["texto"]); d["cuenta"], d["conf"] = r["cuenta"], r["conf"]
        print(f"     {d['numero']:<8} → {r['cuenta'] or '—':<9} ({r['conf']:.2f}) {r['motivo']}")
chk(g["220"]["cuenta"] == "51355001" and g["217"]["cuenta"] == "28150501" and f["cuenta"] == "51953001", "transporte → 51355001 · pago a domicilio → 28150501 · papel → 51953001")
chk(g["216"]["conf"] < 0.75, "«pilas»: cuenta dudosa → la app te pide confirmar")
r = construir("2026-09-12", docs)
for x in r["filas"]: print(f"     {x['codigo']:<9}{x['nombre'][:26]:<27}{x['tercero'][:22]:<23}{x['descripcion'][:34]:<35}{x['debito']:>9,}{x['credito']:>9,}")
chk(sum(1 for d in r["documentos"] if not d["contar"]) == 3, "los 3 recibos manuscritos se detectan como duplicados de los impresos")
chk(r["diferencia"] == 0 and r["debito"] == 320600, f"el día cuadra (la parte con tarjeta, 95.800, también entra): débito {r['debito']:,} = crédito {r['credito']:,}")
caja = [x for x in r["filas"] if x["tercero"] == "Cuantías menores"]
chk(caja[-1]["descripcion"] == "Recibos: 216, 217 y 220" and caja[-1]["credito"] == 114800, "caja menor: «Recibos: 216, 217 y 220» y un crédito a Efectivo por 114.800")
ef = [x for x in r["filas"] if x["tercero"] == "Consumidor Final" and x["medio"] == "Efectivo"]
chk([(x["codigo"], x["debito"], x["credito"]) for x in ef] == [("11050501", 100000, 0), ("28150502", 0, 17800), ("41401501", 0, 82200)], "efectivo 100.000 = propina 17.800 + restaurante 82.200 (la tarjeta cubre primero el consumo)")
tj = [(x["codigo"], x["debito"], x["credito"]) for x in r["filas"] if x["tercero"] == "Consumidor Final" and x["medio"] != "Efectivo"]
chk(tj == [("11200501", 95800, 0), ("41401501", 0, 95800)], f"la parte con tarjeta: Banco 95.800 contra Restaurante: {tj}")

# comandas pagadas con VISA, transferencia sin cifra, Nequi, Rappi y crédito personal (fiado)
def cmd(n, pago, extra=""):
    return "\n".join(["SU RESTAURANTE", f"RECIBO #: {n}", "Fecha de Venta: 16/09/2026 9:00:00 p. m.", extra + "SUBTOTAL CONSUMO $40,000", "**TOTAL CONSUMO>> $50000",
                      "DESCUENTO $10,000", "*TOTAL A PAGAR> $40,000", "PROPINA $", "PAGADO $40,000", pago, "CAMBIO $"])
vv = [leer_texto(cmd(90001, "PAGA CON VISA $40,000"))[0], leer_texto(cmd(90002, "PAGA CON TRANSFEREN"))[0], leer_texto(cmd(90003, "PAGA CON NEQUI $40,0"))[0],
      leer_texto(cmd(90004, "PAGA CON RAPPI $40,000"))[0], leer_texto(cmd(90005, "PAGA CON CRED.PERSONA $40,000", extra="Cliente: ANA PEREZ\n"))[0]]
chk([(x["consumo"], x["tarjeta"], x["transf"], x["credito"], x["medio"]) for x in vv] ==
    [(40000, 40000, 0, 0, "Tarjeta"), (40000, 0, 40000, 0, "Transferencia"), (40000, 0, 40000, 0, "Nequi"), (40000, 0, 40000, 0, "Rappi"), (40000, 0, 0, 40000, "Crédito personal")],
    "VISA = tarjeta · transferencia sin cifra o cortada = lo pagado · Nequi · Rappi · CRED.PERSONA = fiado (no tarjeta) · el descuento no se resta dos veces")
r4 = construir("2026-09-16", vv)
deb = {(x["codigo"], x["medio"]): x["debito"] for x in r4["filas"] if x["debito"]}
chk(deb == {("11200501", "Tarjeta"): 40000, ("11200501", "Transferencia"): 40000, ("11200501", "Nequi"): 40000, ("13452501", "App"): 40000, ("13050501", "Crédito personal"): 40000}
    and [x["credito"] for x in r4["filas"] if x["codigo"] == "41401501"] == [200000] and r4["diferencia"] == 0,
    "tarjeta/transferencia/Nequi → Banco · Rappi → Ingresos por cobrar · fiado → Cuenta por cobrar · Restaurante 200.000")
fx = leer_texto("Factura de venta No. FV123\nCLIENTE: Samu\nTomate 20.000\nTotal Factura: $ 124,400\n")[0]
chk(fx["valor"] == 124400, f"«Total Factura: $ 124,400» → {fx['valor']}")

# comanda mal leída: el efectivo no alcanza para propina + domicilio → no se acredita más de lo que entró
vm = dict(v, id="mal", numero="85599", efectivo=10000, tarjeta=0, domicilio=5000, avisos=[])
r5 = construir("2026-09-12", [vm])
ef5 = [(x["codigo"], x["debito"], x["credito"]) for x in r5["filas"] if x["medio"] == "Efectivo" and (x["debito"] or x["credito"])]
chk(ef5 == [("11050501", 10000, 0), ("28150502", 0, 10000)] and r5["diferencia"] == 0 and any("no alcanza" in a for a in r5["avisos"]),
    f"efectivo 10.000 < propina 17.800 + domicilio 5.000: propina tope 10.000, domicilio 0, el día cuadra y se avisa: {ef5}")
# efectivo 0 (mal leído) con propina y domicilio en efectivo: no hay filas de caja, pero se avisa con las cifras (FIX-4)
r6 = construir("2026-09-12", [dict(vm, id="mal0", efectivo=0)])
chk(not [x for x in r6["filas"] if x["medio"] == "Efectivo" and (x["debito"] or x["credito"])] and r6["diferencia"] == 0
    and any("efectivo es 0" in a and "17.800" in a and "5.000" in a for a in r6["avisos"]),
    f"efectivo 0 con propina 17.800 + domicilio 5.000 en efectivo: sin filas de caja y se avisa: {r6['avisos']}")

# recaudos del banco: las cifras de tu día 12/09
r2 = construir("2026-09-12", [], {"llave": 730550, "tarjetas": 146500, "qr": 1091950})
chk([x["credito"] for x in r2["filas"] if x["credito"]] == [1969000] and r2["diferencia"] == 0, "recaudos del banco: Restaurante 1.969.000 = 730.550 + 146.500 + 1.091.950")
chk(abreviar_recibos(["85531", "85533", "85534", "85540"]) == "85531, 533, 534, 540" and rangos([210, 211, 212, 213]) == "210 al 213" and rangos([208, 209, 214]) == "208, 209 y 214", "frases de descripción como las escribes")
# día que no cuadra
d_mal = leer_texto(GASTO.format(n=300, v="5,000", d="algo raro"))
r3 = construir("2026-09-13", d_mal)
chk(any("sin cuenta" in a for a in r3["avisos"]) and r3["diferencia"] == 0, "un gasto sin cuenta se avisa (no se inventa una)")
# cuentas puente (R10): nada se descarta en silencio y ninguna fila queda sin código
from parser_module.contable.asientos import CFG as CFG_A
PI, PG = CFG_A["puente_ingreso"], CFG_A["puente_gasto"]
chk(r3["filas"] and all(x["codigo"] for x in r3["filas"]) and r3["filas"][0]["codigo"] == PG, f"gasto de caja sin cuenta → {PG} (puente_gasto), ninguna fila sin código")
vs = dict(v, id="sinmedio", numero="85600", efectivo=0, tarjeta=0, transf=0, credito=0, avisos=[])
r6 = construir("2026-09-12", [vs])
p6 = [(x["codigo"], x["debito"], x["credito"]) for x in r6["filas"]]
chk(p6 == [(PI, 178000, 0), ("41401501", 0, 178000)] and r6["diferencia"] == 0 and any("puente" in a for a in r6["avisos"]),
    f"venta sin medio de pago → débito {PI} (puente_ingreso) contra Restaurante y aviso: {p6}")
fs = dict(f, id="sinvalor", valor=0, avisos=[])
fc = dict(f, id="sincuenta", numero="POS1", cuenta="", avisos=[])
r7 = construir("2026-09-12", [fs, fc])
chk([(p["id"], p["valor"]) for p in r7["pendientes"]] == [("sinvalor", 0)] and any("pendiente" in a for a in r7["avisos"]) and r7["diferencia"] == 0,
    "factura sin valor → queda en pendientes con valor 0 y aviso (no se inventa un asiento)")
chk([x["codigo"] for x in r7["filas"]] == [PG, "11200501"], "factura sin cuenta → puente_gasto contra Bancos")
# FIX-1 (R10): un día con solo facturas sin valor se guarda en el libro (los pendientes no quedan fuera)
import tempfile, pathlib, app as appmod
_t = pathlib.Path(tempfile.mkdtemp()); appmod.RUTA_LIBRO, appmod.RUTA_SOPORTE = _t / "libro.json", _t / "soporte.db"   # T15: guardar el día abre soporte.db (R7)
rs = appmod.app.test_client().post("/api/contable/libro", json={"fecha": "2026-09-14", "documentos": [fs]})
gl = appmod._leer_libro().get("2026-09-14", {})
chk(rs.status_code == 200 and [d["id"] for d in gl.get("documentos", [])] == ["sinvalor"] and [p["id"] for p in gl.get("pendientes", [])] == ["sinvalor"],
    f"día con solo una factura sin valor → se guarda con su pendiente ({rs.status_code})")
rv = appmod.app.test_client().post("/api/contable/libro", json={"fecha": "2026-09-15", "documentos": []})
chk(rv.status_code == 400, "día sin papeles ni pendientes → sigue rechazado")
# FIX-6 (R9): el export recalcula; si un día ya no da lo que se confirmó, AVISOS lo dice con las dos cifras
def avisos_export():
    from openpyxl import load_workbook as lw
    w = lw(io.BytesIO(appmod.app.test_client().get("/api/contable/excel").data))
    return [str(c.value or "") for c in w["AVISOS"]["A"]] if "AVISOS" in w.sheetnames else []
chk(not any("cambió" in a for a in avisos_export()), "día recalculado igual a lo confirmado → sin aviso de cambio")
lib_ = appmod._leer_libro(); lib_["2026-09-14"]["debito"] = 999; appmod._guardar_libro(lib_)
av6 = [a for a in avisos_export() if "cambió" in a]
chk(len(av6) == 1 and "2026-09-14" in av6[0] and "999" in av6[0], f"día que cambió desde que se confirmó → aviso con lo guardado y lo de hoy: {av6}")
dias = [r, r3, r4, r5, r6, r7]
contados = [d["id"] for x in dias for d in x["documentos"] if d.get("contar") and d["tipo"] != "cierre"]
en_libro = {p["id"] for x in dias for p in x["pendientes"]} | {i for x in dias for f_ in x["filas"] for i in f_["docs"]}
chk(set(contados) <= en_libro, f"ningún papel desaparece de construir: {sorted(set(contados) - en_libro) or 'todos están'}")
# M-4: una comanda cuyo número no sigue la secuencia del día (85767 entre 85873…) → aviso: ¿de otro día o número mal leído?
sec = [leer_texto(cmd(n, "PAGA CON VISA $40,000"))[0] for n in (85873, 85874, 85876, 85767)]
r6 = construir("2026-09-27", sec)
chk([a.split(":")[0] for a in r6["avisos"] if "secuencia" in a] == ["Comanda 85767"],
    f"comanda 85767 fuera de la secuencia del día (85873–85876) → aviso: {[a for a in r6['avisos'] if 'secuencia' in a]}")
chk(not any("secuencia" in a for a in r4["avisos"]) and not any("secuencia" in a for a in construir("2026-09-27", sec, cfg={"salto_comanda": 200})["avisos"]),
    "comandas seguidas (90001–90005) sin aviso · el salto tolerado se configura (salto_comanda)")
# M-8: mismo número de comanda no es duplicado si la hora de venta o la mesa cambian (el OCR leyó mal el número de otro recibo)
def cmd_hm(n, hora, mesa, iid):
    return dict(leer_texto(cmd(n, "PAGA CON VISA $40,000", extra=f"Mesa: {mesa}\n").replace("9:00:00 p. m.", hora))[0], id=iid)
rep = [cmd_hm(85560, "6:42:34 p. m.", "M1", "a"), cmd_hm(85560, "6:42:34 p. m.", "M30", "b"), cmd_hm(85560, "7:39:13 p. m.", "M1", "c")]
r8 = construir("2026-09-13", rep)
chk([d["contar"] for d in r8["documentos"]] == [True, True, True] and not any(d["dup_de"] for d in r8["documentos"])
    and [a for a in r8["avisos"] if "número repetido" in a] == ["Comanda 85560: número repetido en 3 recibos distintos (no coinciden hora, mesa y valores): revisa el número en el papel"],
    f"85560 con otra mesa u otra hora → los 3 cuentan, con aviso «número repetido: revisa»: {[a for a in r8['avisos'] if 'repetido' in a]}")
copias = [cmd_hm(85505, "8:10:00 p. m.", "M16", i) for i in "pqrst"]   # la misma foto cargada 5 veces (10-09)
r9 = construir("2026-09-10", copias)
chk([d["contar"] for d in r9["documentos"]] == [True, False, False, False, False] and all(d["dup_de"] == "p" for d in r9["documentos"][1:])
    and not any("repetido" in a for a in r9["avisos"]), "misma comanda, misma hora y mesa (la foto cargada 5 veces) → sigue siendo duplicado, sin aviso")
otra = lambda d, **k: dict(d, **{"texto": d["texto"] + "\n(otra lectura)", **k})   # mismo papel leído otra vez: el texto no es idéntico
base = cmd_hm(85580, "7:00:00 p. m.", "M5", "m")
casos = {"otro valor": [base, otra(base, id="n", consumo=52000, tarjeta=52000)],
         "sin mesa": [base, otra(base, id="n", texto=base["texto"].replace("Mesa: M5\n", ""), detalle="")],
         "sin hora": [base, otra(base, id="n", texto=base["texto"].replace("7:00:00 p. m.", ""))]}
for caso, ds in casos.items():
    rc = construir("2026-09-13", ds)
    chk([d["contar"] for d in rc["documentos"]] == [True, True] and any("85580: número repetido en 2" in a for a in rc["avisos"]),
        f"mismo número, misma hora/mesa pero {caso} → son 2 papeles (falta de datos no los junta), con aviso")
par = [cmd_hm(85575, "9:09:18 p. m.", "M40", "u"), otra(cmd_hm(85575, "9:09:18 p. m.", "M40", "w"), consumo=0, tarjeta=0)]   # 13-09: uno leyó 0
rp = construir("2026-09-13", par)
chk([(d["contar"], d["dup_de"]) for d in rp["documentos"]] == [(True, ""), (False, "u")], "85575 13-09: misma hora y mesa, uno con valores en 0 (no leídos) → sigue siendo duplicado")
sn = [dict(cmd_hm(85590, "8:00:00 p. m.", "M7", i), numero="") for i in "ghi"] + [dict(cmd_hm(85591, "8:20:00 p. m.", "M7", "j"), numero="")]
rs = construir("2026-09-10", sn)
chk([d["contar"] for d in rs["documentos"]] == [True, False, False, True] and rs["documentos"][1]["dup_de"] == "g",
    "sin número: la misma foto cargada 3 veces cuenta una vez; otro papel sin número (otra hora) cuenta")
viejo = [cmd_hm(85575, "6:00:00 p. m.", "M2", "x"), dict(cmd_hm(85579, "6:30:00 p. m.", "M4", "y"), dup_de="x", contar=False)]
r10 = construir("2026-09-13", viejo)
chk([(d["dup_de"], d["contar"]) for d in r10["documentos"]] == [("", True), ("", True)], "número corregido (85579): el «duplicado de» viejo se borra y vuelve a contar")
fijo = [cmd_hm(85575, "6:00:00 p. m.", "M2", "x"), dict(cmd_hm(85579, "6:30:00 p. m.", "M4", "y"), dup_de="x", contar=False, manual_contar=True)]
r11 = construir("2026-09-13", fijo)
chk([(d["dup_de"], d["contar"]) for d in r11["documentos"]] == [("", True), ("x", False)], "si tú lo desmarcaste a mano (manual_contar) se respeta tal cual")

# M-9: el OCR marca los papeles con «=== DOCUMENTO ===» pero a veces pega dos recibos en un bloque → se parte también por membrete
pegados = "=== DOCUMENTO ===\n" + cmd(90101, "PAGA CON VISA $40,000") + "\n\n" + cmd(90102, "PAGA CON TRANSFEREN") + "\n=== DOCUMENTO ===\n" + cmd(90103, "EFECTIVO $40,000")
chk([d["numero"] for d in leer_texto(pegados, "comandas")] == ["90101", "90102", "90103"],
    f"M-9: dos recibos pegados en un bloque «DOCUMENTO» salen como dos papeles: {[d['numero'] for d in leer_texto(pegados, 'comandas')]}")

# Excel
x = libro_a_excel({"2026-09-12": r["filas"] + r2["filas"]}, "Mi restaurante", "900.000.000-1")
# guarda de exportación (R10): una fila sin código no llega a MOVIMIENTOS, sale como aviso
from openpyxl import load_workbook
sin_cod = dict(r3["filas"][0], codigo="", descripcion="fila sin código")
wb = load_workbook(io.BytesIO(libro_a_excel({"2026-09-13": [sin_cod]})))
mov = [ws for ws in wb if ws.title.startswith("MOVIMIENTOS")]
n_mov = sum(1 for ws in mov for fila in ws.iter_rows(min_row=9, values_only=True) if any(fila))
av = [c for c in wb["AVISOS"]["A"]] if "AVISOS" in wb.sheetnames else []
chk(n_mov == 0 and any("sin código" in str(c.value) for c in av), f"fila sin código: {n_mov} filas en MOVIMIENTOS y aviso en la hoja AVISOS")
wb2 = load_workbook(io.BytesIO(libro_a_excel({"2026-09-13": r3["filas"]})))
chk("AVISOS" not in wb2.sheetnames, "sin filas malas no hay hoja AVISOS")
# T4 · estado/motivos/lectura/recorte: `_base` usa estado == "revisado" (migra el viejo `revisado`) y todo viaja por el .json crudo v1
from parser_module.contable.asientos import estado_de
from parser_module.contable.exportar import crudo, leer_crudo
chk((estado_de({}), estado_de({"avisos": ["x"]}), estado_de({"avisos": ["x"], "revisado": True}), estado_de({"estado": "dudoso", "revisado": True})) == ("validado", "dudoso", "revisado", "dudoso"),
    "T4: estado se migra al leer (sin estado: validado/dudoso; `revisado` viejo → revisado)")
t4 = dict(leer_texto(VENTA)[0], avisos=["los pagos no cuadran"])
av4 = lambda **k: construir("2026-09-12", [dict(t4, **k)])["avisos"]
chk(any("los pagos no cuadran" in a for a in av4()) and any("los pagos no cuadran" in a for a in av4(estado="dudoso")), "T4: un recibo con avisos sin revisar los muestra")
chk(not any("los pagos no cuadran" in a for a in av4(estado="revisado")) and not any("los pagos no cuadran" in a for a in av4(revisado=True)), "T4: estado «revisado» (o el viejo revisado) silencia sus avisos")
cand = lambda v, var, m: dict(valor=v, variante=var, motor="qwen2.5vl")
lec = dict(pasadas=3, candidatos={c: [cand(1000, "V0", 0), cand(1000, "V1", 0), cand(900, "V2", 0)] for c in ("consumo", "propina", "efectivo", "tarjeta", "numero")},
           nitidez=0.8, tenengrad=123.5, segundos=14.2)
t4l = dict(t4, id="2026-09-12/t4", estado="dudoso", motivos=["Los medios de pago suman 51.500 y pagado − cambio es 52.000"], lectura=lec, recorte="ab12cd34.jpg")
lib4 = {"2026-09-12": dict(fecha="2026-09-12", documentos=[t4l], banco={}, cierre={}, guardado="2026-09-12T08:00:00")}
d4, m4 = leer_crudo(crudo(lib4))
chk(not m4 and d4[0]["documentos"][0] == t4l, "T4: documento con lectura completa (candidatos, motivos, recorte) hace la ida y vuelta por el .json crudo v1")
import tempfile, os; open(os.path.join(tempfile.gettempdir(), "libro_prueba.xlsx"), "wb").write(x)
# M-10: la fecha de la carpeta («20 SEPT 26») manda sobre la impresa; el nombre del archivo trae la fecha del escaneo y no cuenta
from parser_module.contable.documentos import fecha_de_carpeta, aplicar_fecha_carpeta
R = "SEPTIEMBRE/SEPTIEMBRE DEL 15 AL 30/20 SEPT 26/CIERRE/cierre20261002_18131736.pdf"
chk(fecha_de_carpeta(R) == "2026-09-20", f"M-10: «20 SEPT 26» → 2026-09-20, no la del nombre (escaneo 02-10): {fecha_de_carpeta(R)}")
chk(fecha_de_carpeta("x/17 sept 26/COMANDAS/a.pdf") == "2026-09-17" and fecha_de_carpeta(r"x\3 de octubre de 2026\a.jpg") == "2026-10-03",
    "M-10: minúsculas, «de», año de 4 cifras y rutas de Windows")
chk(fecha_de_carpeta("comandas20260920_85735.pdf") is None and fecha_de_carpeta("x/31 SEPT 26/a.jpg") is None
    and fecha_de_carpeta("6/comandas/a.jpg", 2026) is None and fecha_de_carpeta("SEPTIEMBRE/31/comandas/a.jpg", 2026) is None,
    "M-10: fecha solo en el nombre, fecha imposible o día sin mes en una carpeta de arriba → None (sigue la impresa)")
chk(fecha_de_carpeta("SEPTIEMBRE/6/comandas/85447 al 452.jpg", 2026) == "2026-09-06" and fecha_de_carpeta(r"D:\data\SEPTIEMBRE\11\comandas\10-09.jpg", 2026) == "2026-09-11"
    and fecha_de_carpeta("SEPTIEMBRE/6/comandas/a.jpg") == f"{__import__('datetime').date.today().year}-09-06",
    "M-10: carpeta de solo el día → día de la carpeta + mes de la carpeta de arriba + año dado (sin año: el de hoy); el nombre «10-09» no cuenta")
dm = [dict(tipo="venta", numero="85767", fecha="2026-09-27", avisos=[]), dict(tipo="venta", numero="85768", fecha="2026-09-20", avisos=[]),
      dict(tipo="venta", numero="85769", fecha="", avisos=[])]
aplicar_fecha_carpeta(dm, "S/20 SEPT 26/COMANDAS/comandas20260920_85766 al 768.pdf")
chk([d["fecha"] for d in dm] == ["2026-09-20"] * 3 and dm[0].get("fecha_impresa") == "2026-09-27" and any("27/09/2026" in a for a in dm[0]["avisos"])
    and not dm[1]["avisos"] and not dm[2]["avisos"] and "fecha_impresa" not in dm[1], "M-10: 85767 se queda el 20-09 con aviso de la fecha impresa; sin aviso si coincide o no se leyó")
d0 = [dict(tipo="venta", fecha=f, avisos=[]) for f in ("2026-09-06", "2025-09-06", "2026-09-05", "")]
aplicar_fecha_carpeta(d0, "SEPTIEMBRE/6/comandas/85447 al 452.jpg")
chk([d["fecha"] for d in d0] == ["2026-09-06"] * 4 and [bool(d["avisos"]) for d in d0] == [False, True, True, False] and d0[1]["fecha_impresa"] == "2025-09-06",
    "M-10: «SEPTIEMBRE/6/» con el año más común de las fechas impresas (2026): el «2025» de M-12 queda 2026-09-06 con aviso")
d1 = [dict(tipo="venta", fecha="2025-09-06", avisos=[])]; aplicar_fecha_carpeta(d1, "6/comandas/a.jpg")
chk(d1[0]["fecha"] == "2025-09-06" and not d1[0]["avisos"], "M-10: sin mes en la ruta no toca nada")
# M-13 T2 · fitness ADR-008: un documento marcador tipo "faltante" no mueve nada
from parser_module.contable import faltantes as _falt
from parser_module.contable.validar import faltantes as _validar_falt
_base_t2 = [leer_texto(VENTA)[0], leer_texto(GASTO)[0]]   # mismos ids en las dos corridas
def _dia(extra):
    return construir("2026-09-12", [dict(d) for d in _base_t2] + extra)
_sin, _con = _dia([]), _dia(_falt.crear([85999], "foto.jpg", "2026-09-12", "ab12cd34.jpg"))
chk(any(d["tipo"] == "faltante" for d in _con["documentos"]), "T2: el marcador viaja en el día")
chk(_sin["filas"] == _con["filas"] and _sin["datos"] == _con["datos"] and (_sin["debito"], _sin["credito"]) == (_con["debito"], _con["credito"]),
    "ADR-008: un faltante movió los asientos")
chk(_sin["avisos"] == _con["avisos"] and _sin["pendientes"] == _con["pendientes"] and _sin["cruces"] == _con["cruces"] and _sin["faltan"] == _con["faltan"],
    "ADR-008: un faltante movió los avisos, pendientes o el cruce con el cierre")
chk(not any(x["tipo"] == "faltante" for x in _con["datos"]), "ADR-008: el marcador no sale como fila de DATOS OCR (Excel)")
_marca = _falt.crear([5], "f.jpg", "2026-09-12", "r.jpg")
_ven = [dict(tipo="venta", numero="5")]
chk(_validar_falt(_ven + _marca, [5, 6]) == [6] and _validar_falt(_marca, [5]) == [5], "ADR-008: un marcador cuenta como recibo leído")
_lib = {"2026-09-12": dict(fecha="2026-09-12", documentos=_con["documentos"], banco={}, cierre={}, guardado="2026-09-12T08:00:00")}
_d, _m = leer_crudo(crudo(_lib))
chk(not _m and _d[0]["documentos"] == _con["documentos"], "T2: el marcador hace la ida y vuelta por el .json crudo")
# R1 (v1.3.0): ajuste al cierre cuando el cierre supera la venta de los papeles
def _r1(cierre=None, extra=()):
    return construir("2026-09-16", [dict(d) for d in vv[:2]] + list(extra), None, None, cierre)   # 2 comandas de 40.000 = 80.000
_a = _r1({"total_vendido": 200000})
_aj = [(x["codigo"], x["debito"], x["credito"]) for x in _a["filas"] if x["descripcion"] == "Ajuste al cierre (cierre > papeles)"]
_rest = sum(x["credito"] for x in _a["filas"] if x["codigo"] == "41401501")
chk(_aj == [("13809501", 120000, 0), ("41401501", 0, 120000)] and _rest == 200000 and any("Ajuste al cierre" in t for t in _a["avisos"]),
    f"R1: cierre 200.000 vs papeles 80.000 → 4140 suma 200.000, puente 120.000 y aviso: {_aj}")
_cd = dict(tipo="cierre", id="c1", fecha="2026-09-16", contar=False, datos={"total_vendido": 200000}, numero="", tercero="POS", detalle="Cierre de caja resumido", avisos=[], texto="", origen="")
_b = _r1(extra=[_cd])
chk([(x["codigo"], x["debito"], x["credito"]) for x in _b["filas"] if x["descripcion"].startswith("Ajuste")] == [(x["codigo"], x["debito"], x["credito"]) for x in _a["filas"] if x["descripcion"].startswith("Ajuste")]
    and _b["diferencia"] == 0 and _a["diferencia"] == 0 and all(x["fx"] for x in _a["filas"] if x["descripcion"].startswith("Ajuste")),
    "R1: igual si el total sale del papel del cierre (con su fecha); débito = crédito; cada fila nueva lleva fx")
_m = _r1({"total_vendido": 50000}); _s = _r1()
chk(_m["filas"] == _s["filas"] and any("más que el cierre" in t for t in _m["avisos"]) and not any("Ajuste" in t for t in _s["avisos"]) and _m["diferencia"] == 0,
    "R1: papeles > cierre → filas iguales que sin cierre + aviso; sin cierre → sin ajuste ni aviso")
_o = _r1(extra=[dict(_cd, fecha="2026-09-17")])
chk(_o["filas"] == _s["filas"], "R1: cierre de otra fecha → sin ajuste")
# R2 (v1.3.0): débito a 2815 desde el cierre, sin repetir los vales «pago domicilios»
_p = _r1({"total_vendido": 80000, "propinas_total": 48100})
_pd = [(x["codigo"], x["debito"], x["credito"]) for x in _p["filas"] if "según cierre" in x["descripcion"]]
chk(_pd == [("28150502", 48100, 0), ("11050501", 0, 48100)] and _p["diferencia"] == 0,
    f"R2: propinas 48.100 en el cierre → débito 28150502 48.100 contra Efectivo: {_pd}")
_vale = dict(g["217"], id="vale20", valor=20000, cuenta="28150501", contar=True, fecha="2026-09-16", avisos=[])
_q = _r1({"total_vendido": 80000, "domicilios_pagados": 50000}, extra=[_vale])
_qd = [x["debito"] for x in _q["filas"] if x["codigo"] == "28150501" and x["debito"]]
_qc = [x["debito"] for x in _q["filas"] if x["codigo"] == "28150501" and "según cierre" in x["descripcion"]]
chk(sorted(_qd) == [20000, 30000] and _qc == [30000] and _q["diferencia"] == 0,
    f"R2: vale 20.000 + cierre 50.000 → 30.000 desde el cierre, 28150501 suma 50.000 una sola vez: {_qd}")
_q2 = _r1({"total_vendido": 80000, "domicilios_pagados": 10000}, extra=[_vale])
chk(not any(x["codigo"] == "28150501" and "según cierre" in x["descripcion"] for x in _q2["filas"]) and _q2["diferencia"] == 0
    and all(x["fx"] for x in _p["filas"] + _q["filas"] if "según cierre" in x["descripcion"]) and not any("según cierre" in x["descripcion"] for x in _s["filas"]),
    "R2: vales > cierre → no hay fila desde el cierre; débito = crédito; cada fila nueva lleva fx; sin cierre → sin filas nuevas")
# R3 (v1.3.0): respaldo único del libro antes de la primera escritura con las reglas nuevas
import json
_ant = appmod.RUTA_LIBRO; _t3 = pathlib.Path(tempfile.mkdtemp()); appmod.RUTA_LIBRO = _t3 / "libro.json"
_orig = {"2026-09-16": dict(fecha="2026-09-16", documentos=[dict(d) for d in vv[:2]], banco={}, cierre={"total_vendido": 200000}, guardado="2026-09-16T08:00:00")}
_sin = construir("2026-09-16", [dict(d) for d in vv[:2]], None, None, None)
_orig["2026-09-16"].update(debito=_sin["debito"], credito=_sin["credito"])   # confirmado antes de R1: sin ajuste
appmod.RUTA_LIBRO.write_text(json.dumps(_orig), encoding="utf8"); _txt0 = appmod.RUTA_LIBRO.read_text(encoding="utf8")
_bk = _t3 / "libro.antes-v1.3.0.json"
_nuevo = dict(_orig); _nuevo["2026-09-17"] = dict(_orig["2026-09-16"], fecha="2026-09-17")
appmod._guardar_libro(_nuevo)
chk(_bk.exists() and _bk.read_text(encoding="utf8") == _txt0, "R3: la primera escritura crea el respaldo idéntico al libro original")
_bk.write_text("marca", encoding="utf8"); appmod._guardar_libro(_orig)
chk(_bk.read_text(encoding="utf8") == "marca", "R3: la segunda escritura no toca el respaldo")
appmod._guardar_libro(_orig)
_cam = [a for a in avisos_export() if "cambió" in a]
chk(len(_cam) == 1 and "2026-09-16" in _cam[0], f"R3: día con cierre > papeles confirmado antes → el Excel lo lista en «cambió desde que se confirmó»: {_cam}")
appmod.RUTA_LIBRO = _ant
# --- Rappi (01c #6) T1: el papel de Rappi trae ID largo y «Total sugerido a facturar»
PAPEL_RAPPI = """Rappi
Juan Perez
ID 1000000001
Restaurante: Su Restaurante
Plataforma: WEB
Fecha orden: 2026-09-02 21:00
1 Bandeja paisa | 1 | $25.000
Subtotal $25.000
Total sugerido a facturar $25.000
VERIFICA LOS PRODUCTOS
Juan Perez
ID 1000000001"""
_rp = leer_texto(PAPEL_RAPPI, "comandas")[0]
chk(_rp["tipo"] == "venta" and _rp.get("rappi") and _rp["numero"] == "1000000001" and _rp["valor"] == 25000 and _rp["consumo"] == 25000 and _rp["medio"] == "Rappi" and _rp["tercero"] == "Rappi" and not _rp["avisos"],
    f"Rappi: papel real -> ID 1000000001, valor 25.000, medio Rappi: {_rp['numero']} {_rp['valor']} {_rp['medio']} {_rp['avisos']}")
_rr = construir("2026-09-17", [_rp])
_f = [f for f in _rr["filas"] if f["codigo"] == "13452501"]
chk(len(_f) == 1 and _f[0]["debito"] == 25000 and _f[0]["tercero"] == "Rappi" and _f[0]["medio"] == "App" and _rr["debito"] == _rr["credito"],
    f"Rappi: construir -> 13452501 debito 25.000 tercero Rappi medio App, debito = credito: {[(f['codigo'], f['debito'], f['tercero'], f['medio']) for f in _rr['filas']]}")
_rs = leer_texto(PAPEL_RAPPI.replace("Total sugerido a facturar $25.000", "").replace("Subtotal $25.000", ""), "comandas")[0]
chk(_rs["valor"] == 0 and any("total de Rappi" in a for a in _rs["avisos"]), f"Rappi: papel sin total -> valor 0 + aviso: {_rs['valor']} {_rs['avisos']}")
print("\n" + ("TODO OK" if ok else "HAY FALLOS"))
