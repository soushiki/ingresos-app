"""Prueba de la hoja DATOS OCR (R9, ADR-002): una fila por papel y medio de pago; ningún papel falta, ni los que van a puente."""
import copy, io, json, re, zipfile
from pathlib import Path
from openpyxl import load_workbook
from parser_module.contable import construir, leer_texto, libro_a_excel
from parser_module.contable.exportar import MESES
MESES_ = {m: i for i, m in enumerate(MESES, 1)}
from parser_module.contable.asientos import CFG, COLS_DATOS
from parser_module.contable.ejemplos import FOTO_PAPELES, GASTO, VENTA
ok = True
def chk(c, m):
    global ok; ok &= bool(c); print(("✓ " if c else "✗ ") + m)

PI, PG = CFG["puente_ingreso"], CFG["puente_gasto"]
docs = leer_texto(FOTO_PAPELES) + leer_texto(VENTA)
v = docs[-1]; f = next(d for d in docs if d["tipo"] == "factura")
docs += [dict(v, id="sinmedio", numero="85600", efectivo=0, tarjeta=0, transf=0, credito=0, avisos=[]), dict(f, id="sinvalor", numero="POS2", valor=0, avisos=[])]
docs += [dict(v, id="ilegible", numero="85601", consumo=0, propina=0, efectivo=0, tarjeta=0, transf=0, credito=0, avisos=[])]
docs += leer_texto(GASTO.format(n=300, v="5,000", d="algo raro"))   # gasto de caja sin cuenta
r = construir("2026-09-12", docs, {"tarjetas": 146500})
D = r["datos"]
chk(set(D[0]) == set(COLS_DATOS), f"columnas de ADR-002: {COLS_DATOS}")
ids = [x["id"] for x in D if x["id"]]
chk(set(ids) == {d["id"] for d in r["documentos"]}, f"cada papel del día tiene su fila: {len(set(ids))} de {len(r['documentos'])}")
p = {x["id"]: x for x in D}
v_ = [x for x in D if x["id"] == v["id"]]
chk([(x["medio"], x["valor"], x["cuenta"], x["estado"]) for x in v_] == [("Efectivo", 100000, "11050501", "ok"), ("Tarjeta", 95800, "11200501", "en extracto")],
    f"venta mixta → 2 filas (efectivo y tarjeta; la tarjeta la cubre el extracto): {[(x['medio'], x['valor'], x['estado']) for x in v_]}")
chk(p["sinmedio"]["cuenta"] == PI and p["sinmedio"]["medio"] == "(sin medio)" and p["sinmedio"]["aviso"], f"venta sin medio → fila en {PI} con aviso")
g = next(x for x in D if x["numero"] == "300")
chk(g["cuenta"] == PG and g["estado"] == "ok" and g["aviso"], f"gasto sin cuenta → fila en {PG} con aviso")
chk(p["ilegible"]["estado"] == "pendiente" and p["ilegible"]["aviso"], "venta sin ninguna cifra leída → fila «pendiente» (no se pierde de vista)")
chk(p["sinvalor"]["estado"] == "pendiente" and p["sinvalor"]["valor"] == 0, "factura sin valor → estado «pendiente»")
dup = [x for x in D if x["estado"] == "no cuenta"]
chk(len(dup) == 3 and all("duplicado" in x["aviso"] for x in dup), f"los 3 manuscritos duplicados se ven como «no cuenta»: {len(dup)}")
ext = [x for x in D if x["tipo"] == "extracto"]
chk([(x["valor"], x["medio"], x["bloque"]) for x in ext] == [(146500, "Tarjeta", "extracto")], "el extracto bancario también es una fila")
# los importes que suman («ok») dan lo mismo que el libro, por bloque
caja = sum(x["valor"] for x in D if x["bloque"] == "caja" and x["estado"] == "ok")
chk(caja == next(x["credito"] for x in r["filas"] if x["tercero"] == "Cuantías menores" and x["credito"]), f"caja menor en DATOS OCR = crédito a Efectivo del libro ({caja:,})")

# Excel: DATOS OCR va primero, con una fila por dato y sin macros
x = libro_a_excel({"2026-09-12": r["filas"]}, datos={"2026-09-12": D})
wb = load_workbook(io.BytesIO(x))
ws = wb.worksheets[0]
filas = [f_ for f_ in ws.iter_rows(min_row=2, values_only=True) if any(f_)]
chk(ws.title == "DATOS OCR" and len(filas) == len(D), f"hoja DATOS OCR primera, {len(filas)} filas = {len(D)} datos")
chk(not any("vbaProject" in n for n in zipfile.ZipFile(io.BytesIO(x)).namelist()), "ADR-002: el libro no trae macros")
chk("DATOS OCR" not in load_workbook(io.BytesIO(libro_a_excel({"2026-09-12": r["filas"]}))).sheetnames, "sin datos no se agrega la hoja (compatibilidad)")


# T6 · MOVIMIENTOS con fórmulas. openpyxl no calcula y aquí no hay Excel: un evaluador mínimo (refs, SUMIFS, MIN, SUM, + y −) hace de Excel.
TOK = re.compile(r"\s*(?:(?P<num>\d+(?:\.\d+)?)|(?P<str>\"(?:[^\"]|\"\")*\")|(?P<ref>(?:'[^']+'!)?\$?[A-Z]+\$?\d*(?::\$?[A-Z]+\$?\d*)?)(?P<fn>\()?|(?P<op>[-+(),=]))")
def evaluar(wb, hoja, celda, memo):
    if (hoja, celda) in memo: return memo[hoja, celda]
    v = wb[hoja][celda].value
    if isinstance(v, str) and v.startswith("="):
        toks = [m for m in TOK.finditer(v[1:]) if m.group()]; i = 0
        def rango(t):
            h, _, r_ = t.rpartition("!"); h = h.strip("'") or hoja; a, _, b = r_.replace("$", "").partition(":")
            ws = wb[h]; ca, cb = re.sub(r"\d", "", a), re.sub(r"\d", "", b)
            fa, fb = (int(re.sub(r"\D", "", a)), int(re.sub(r"\D", "", b))) if re.search(r"\d", a) else (1, ws.max_row)
            return [[(h, f"{c}{k}") for k in range(fa, fb + 1)] for c in sorted({ca, cb})][0]
        def igual(x, crit):
            if isinstance(crit, str):
                neg = crit.startswith("<>"); c = crit[2:] if neg else crit.removeprefix("=")
                c = re.sub(r"~(.)", r"\1", c); s = "" if x is None else str(x)
                e = s.lower() == c.lower() or (s.replace(".0", "") == c)
                return e != neg
            return x == crit
        def atomo():
            nonlocal i
            t = toks[i]; i += 1
            if t["op"] == "(": x = expr(); i += 1; return x
            if t["num"]: return float(t["num"])
            if t["str"]: return t["str"][1:-1].replace('""', '"')
            if t["fn"]:
                args = []
                while True:
                    if toks[i]["ref"] and ":" in toks[i]["ref"] and toks[i + 1]["op"] in (",", ")"): args.append(rango(toks[i]["ref"])); i += 1
                    else: args.append(expr())
                    i += 1
                    if toks[i - 1]["op"] == ")": break
                fn = t["ref"].upper()
                if fn == "MIN": return min(args)
                if fn == "IF": return args[1] if args[0] else args[2]
                if fn == "SUM": return sum(evaluar(wb, h, c, memo) or 0 for h, c in args[0])
                if fn == "SUMIFS":
                    tot = 0
                    for k, (h, c) in enumerate(args[0]):
                        if k == 0: continue   # encabezado
                        if all(igual(wb[args[n][k][0]][args[n][k][1]].value, args[n + 1]) for n in range(1, len(args), 2)): tot += wb[h][c].value or 0
                    return tot
                raise ValueError(fn)
            h, _, c = t["ref"].rpartition("!")
            return evaluar(wb, h.strip("'") or hoja, c.replace("$", ""), memo)
        def suma():
            nonlocal i
            x = atomo()
            while i < len(toks) and toks[i]["op"] in ("+", "-"):
                o = toks[i]["op"]; i += 1; y = atomo(); x = x + y if o == "+" else x - y
            return x
        def expr():   # comparación «=» (para IF), sin distinguir mayúsculas como Excel
            nonlocal i
            x = suma()
            if i < len(toks) and toks[i]["op"] == "=": i += 1; y = suma(); return igual(x, y)
            return x
        v = expr()
    memo[hoja, celda] = v; return v

def movimientos(wb):
    ws = next(w for w in wb.worksheets if w.title.startswith("MOVIMIENTOS")); memo = {}
    return ws, [(r, evaluar(wb, ws.title, f"F{r}", memo), evaluar(wb, ws.title, f"G{r}", memo)) for r in range(9, ws.max_row + 1) if ws[f"B{r}"].value]

ws, mov = movimientos(wb)
imp = [ws[f"{'G' if f['fx'][0] == 'credito' else 'F'}{k}"].value for k, f in enumerate(r["filas"], 9)]
chk(all(isinstance(x, str) and x.startswith("=") for x in imp), f"ADR-002: cada importe de MOVIMIENTOS es una fórmula ({len(imp)} filas)")
chk(any("'DATOS OCR'!" in x and "SUMIFS" in x for x in imp) and any(re.fullmatch(r"=IF\('DATOS OCR'!J(\d+)=\"ok\",'DATOS OCR'!G\1,0\)", x) for x in imp), "hay SUMIFS por fecha/estado/cuenta y referencias a DATOS OCR que miran el estado")
chk([(int(d), int(c)) for _, d, c in mov] == [(f["debito"], f["credito"]) for f in r["filas"]], "evaluadas, las fórmulas dan lo mismo que construir()")
# corregir a mano en DATOS OCR: el efectivo de la venta mixta pasa de 100.000 a 130.000
fe = next(k for k, x in enumerate(D, 2) if x["id"] == v["id"] and x["medio"] == "Efectivo")
wb["DATOS OCR"][f"G{fe}"] = 130000
_, mov2 = movimientos(wb)
cambio = {r_: (d2 - d1, c2 - c1) for (r_, d1, c1), (_, d2, c2) in zip(mov, mov2) if (d1, c1) != (d2, c2)}
fil = {k: f for k, f in enumerate(r["filas"], 9)}
chk(sorted((fil[k]["codigo"], dc) for k, dc in cambio.items()) == [("11050501", (30000, 0)), ("41401501", (0, 30000))],
    f"cambiar un valor en DATOS OCR mueve Efectivo y la contrapartida de ingresos: {[(fil[k]['codigo'], dc) for k, dc in cambio.items()]}")
chk(evaluar(wb, ws.title, "F7", {}) == evaluar(wb, ws.title, "G7", {}), "y el mes sigue cuadrando (total débito = total crédito)")
wb["DATOS OCR"][f"J{fe}"] = "no cuenta"   # marcar el papel como «no cuenta» lo saca del libro
_, mov3 = movimientos(wb)
chk(sum(d for _, d, _ in mov3) == sum(f["debito"] for f in r["filas"]) - 100000, "cambiar el estado a «no cuenta» saca ese valor del libro")
# FIX-2: las filas «dato» (facturas, extracto, Rappi) también respetan el estado
kf = next(k for k, f_ in enumerate(r["filas"], 9) if f_.get("fx", [None, [None]])[1][0] == "dato" and f_.get("docs") == [f["id"]])
ff = next(k for k, x_ in enumerate(D, 2) if x_["id"] == f["id"] and x_["valor"])
wb["DATOS OCR"][f"J{ff}"] = "no cuenta"
_, mov4 = movimientos(wb)
chk(dict((q, d) for q, d, _ in mov4)[kf] == 0, f"una factura en «no cuenta» → 0 en MOVIMIENTOS (fila {kf}: {dict((q, d) for q, d, _ in mov4)[kf]:,})")
chk(evaluar(wb, ws.title, "F7", {}) == evaluar(wb, ws.title, "G7", {}), "y el mes sigue cuadrando sin esa factura")
av = load_workbook(io.BytesIO(x))["AVISOS"]
chk(any("grupo nuevo" in (c.value or "") for c in av["A"]), "AVISOS explica que un grupo nuevo no agrega fila en MOVIMIENTOS")

# fitness ADR-002 con los días guardados (si hay libro.json): cada día evaluado = construir()
LIB = Path(__file__).parent / "datos" / "libro.json"
if LIB.exists():
    lib = json.loads(LIB.read_text(encoding="utf-8")); rs = {k: construir(k, copy.deepcopy(d["documentos"]), d.get("banco"), None, d.get("cierre")) for k, d in lib.items()}
    wb = load_workbook(io.BytesIO(libro_a_excel({k: q["filas"] for k, q in rs.items()}, datos={k: q["datos"] for k, q in rs.items()})))
    malos, n = [], 0
    for w in [w for w in wb.worksheets if w.title.startswith("MOVIMIENTOS")]:
        memo, m = {}, MESES_[w.title.split()[1]]; esp = [(f["debito"], f["credito"]) for k in sorted(rs) if int(k[5:7]) == m for f in rs[k]["filas"]]
        got = [(int(evaluar(wb, w.title, f"F{q}", memo)), int(evaluar(wb, w.title, f"G{q}", memo))) for q in range(9, w.max_row + 1) if w[f"B{q}"].value]
        n += len(got); malos += [q for q, (a, b) in enumerate(zip(got, esp)) if a != b] + ([-1] if len(got) != len(esp) else [])
    chk(not malos, f"libro.json: {len(rs)} días, {n} asientos con fórmula = construir() (distintos: {len(malos)})")
print("\n" + ("TODO OK" if ok else "HAY FALLOS"))
