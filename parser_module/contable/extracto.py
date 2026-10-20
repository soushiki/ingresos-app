"""Extracto bancario (01c #5). Reglas y clasificacion: concepto y destino salen de extracto_reglas.json (ADR-013).
Ningun codigo de cuenta vive aqui: las cuentas estan en el JSON con su fuente y la marca por_confirmar (regulado).
"""
import datetime
import hashlib
import io
import json
import re
import unicodedata
from pathlib import Path

from .asientos import PUENTES, RECAUDOS_BANCO

RUTA_REGLAS = Path(__file__).with_name("extracto_reglas.json")
RUTA_PLAN = Path(__file__).with_name("plan_de_cuentas.json")
SIGNOS = ("+", "-", "±")
SIN_CONCEPTO = "SIN CONCEPTO"
CASILLAS = {k for k, _, _ in RECAUDOS_BANCO} | {"rappi"}


MAX_BYTES = 5 * 1024 * 1024
MAX_FILAS = 5000
HOJA = "Hoja 1"
COLUMNAS = ("FECHA", "DESCRIPCION", "REFERENCIA", "VALOR")


class ExtractoInvalido(ValueError):
    """El archivo del banco no sirve: el mensaje dice por que (para mostrarlo tal cual)."""


class ReglasInvalidas(ValueError):
    """extracto_reglas.json mal editado: el mensaje dice cual regla y por que."""


def normalizar(txt) -> str:
    "Mayusculas, sin tildes y espacios colapsados (sin espacios al inicio; uno final si el texto lo trae)."
    t = unicodedata.normalize("NFKD", str(txt or "")).encode("ascii", "ignore").decode().upper()
    return re.sub(r"\s+", " ", t.lstrip())


def _signo(s) -> str:
    s = str(s).strip()
    return "-" if s in ("−", "–") else s


def cargar_reglas(ruta=None) -> dict:
    "Lee y valida las reglas; lanza ReglasInvalidas con el motivo si algo esta roto."
    ruta = Path(ruta) if ruta else RUTA_REGLAS
    try:
        r = json.loads(ruta.read_text(encoding="utf8"))
    except (OSError, ValueError) as e:
        raise ReglasInvalidas(f"no se pudo leer {ruta.name}: {e}") from e
    if not isinstance(r, dict):
        raise ReglasInvalidas("el JSON debe ser un objeto")
    plan = {x["codigo"] for x in json.loads(RUTA_PLAN.read_text(encoding="utf8"))} | set(PUENTES)
    cuentas = r.get("cuentas")
    if not isinstance(cuentas, dict) or not cuentas:
        raise ReglasInvalidas("falta «cuentas»")
    for k, c in cuentas.items():
        if not isinstance(c, dict) or not re.fullmatch(r"[0-9]{8}", str(c.get("codigo", ""))):
            raise ReglasInvalidas(f"cuenta «{k}»: el codigo debe tener 8 digitos")
        if c["codigo"] not in plan:
            raise ReglasInvalidas(f"cuenta «{k}»: {c['codigo']} no esta en el plan de cuentas ni en los puentes")
        if not str(c.get("fuente", "")).strip() or not isinstance(c.get("por_confirmar"), bool):
            raise ReglasInvalidas(f"cuenta «{k}»: faltan «fuente» o «por_confirmar» (true/false)")
    v = r.get("ventana_factura_dias")
    if not isinstance(v, int) or isinstance(v, bool) or v < 1:
        raise ReglasInvalidas("«ventana_factura_dias» debe ser un entero >= 1")
    st = r.get("sin_tercero")
    if not isinstance(st, list) or any(not isinstance(p, str) or not p.strip() for p in st):
        raise ReglasInvalidas("«sin_tercero» debe ser una lista de prefijos no vacios")
    pt = r.get("prefijos_tercero")
    if not isinstance(pt, list) or any(not isinstance(p, str) or not p.strip() for p in pt):
        raise ReglasInvalidas("«prefijos_tercero» debe ser una lista de prefijos no vacios")
    reglas = r.get("reglas")
    if not isinstance(reglas, list) or not reglas:
        raise ReglasInvalidas("falta la lista «reglas»")
    for i, g in enumerate(reglas, 1):
        n = f"regla {i}"
        if not isinstance(g, dict):
            raise ReglasInvalidas(f"{n}: debe ser un objeto")
        if not isinstance(g.get("prefijo"), str) or not g["prefijo"].strip():
            raise ReglasInvalidas(f"{n}: prefijo vacio")
        if _signo(g.get("signo")) not in SIGNOS:
            raise ReglasInvalidas(f"{n} ({g['prefijo']!r}): signo debe ser +, - o ±")
        if not str(g.get("concepto", "")).strip():
            raise ReglasInvalidas(f"{n} ({g['prefijo']!r}): concepto vacio")
        d = g.get("destino")
        tipo, _, clave = str(d).partition(":")
        if not (d == "revisar" or (tipo == "casilla" and clave in CASILLAS) or (tipo == "cuenta" and clave in cuentas)):
            raise ReglasInvalidas(f"{n} ({g['prefijo']!r}): destino desconocido {d!r}")
    return r


AV_CONFIRMAR, AV_CUAL = "enlazada a una factura pagada por banco: confirma", "¿cuál factura?"
_AVISOS_CLASIF = ("signo inesperado", "sin concepto: ninguna regla encaja", AV_CONFIRMAR, AV_CUAL)


def _dia(f):
    try:
        return datetime.date.fromisoformat(str(f)[:10])
    except ValueError:
        return None


def candidatas(linea: dict, facturas: list, ventana: int, usadas=()) -> list:
    "ADR-015: documentos de banco (no efectivo) con el mismo valor exacto que la salida y fecha a +-ventana dias, aun sin enlazar."
    dl, v = _dia(linea.get("fecha")), -(linea.get("valor") or 0)
    if dl is None or v <= 0:
        return []
    return [f for f in facturas if f.get("tipo") in ("factura", "gasto_caja", "caja_menor_manuscrito") and not f.get("dup_de")
            and f.get("medio") not in ("", None, "Efectivo") and f.get("valor") == v and f.get("id") not in usadas
            and (df := _dia(f.get("fecha"))) is not None and abs((dl - df).days) <= ventana]


def tercero_de(descripcion, reglas: dict) -> str:
    "Nombre normalizado del tercero (sin prefijo ni numeros sueltos); '' si la descripcion es opaca (sin_tercero) o no trae nombre."
    d = normalizar(descripcion).strip()
    if any(d.startswith(normalizar(p).strip()) for p in reglas.get("sin_tercero", [])):
        return ""
    for p in reglas.get("prefijos_tercero", []):
        p = normalizar(p)
        if d.startswith(p.strip()) and (d == p.strip() or d[len(p.strip()):len(p.strip()) + 1] == " "):
            resto = d[len(p.strip()):]
            return " ".join(t for t in resto.split() if not t.isdigit())
    return ""


def _cuenta_recordada(memoria: dict, tercero: str) -> str:
    "Codigo de 8 digitos recordado para el tercero (valor texto o {cuenta}); '' si no hay o esta mal formado."
    m = (memoria or {}).get(tercero)
    c = m.get("cuenta") if isinstance(m, dict) else m
    return c if isinstance(c, str) and re.fullmatch(r"[0-9]{8}", c) else ""


def clasificar(lineas: list, reglas: dict, memoria: dict = {}, facturas: list = [], usadas=()) -> list:
    """Pura (FIX-3: `usadas` = factura_ids ya enlazadas por lineas de otros dias, no se ofrecen): devuelve copias de las lineas con concepto, destino, cuenta, origen y avisos. Memoria por tercero (T7).
    Facturas (T8, ADR-015): una salida que iba a revisar y tiene UNA factura de banco candidata -> destino factura + factura_id (sin asiento, a confirmar);
    con 2 o mas -> revisar con «¿cuál factura?» y sin_asiento (no hay puente). Una factura no se enlaza a dos lineas."""
    out, usadas = [], set(usadas)
    for ln in lineas:
        # los avisos de clasificacion se recalculan desde cero (idempotente); otros avisos se conservan
        x = dict(ln, avisos=[a for a in (ln.get("avisos") or []) if a not in _AVISOS_CLASIF])
        desc, valor = normalizar(x.get("descripcion")), x.get("valor") or 0
        signo = "+" if valor > 0 else "-" if valor < 0 else ""
        misma, hit = False, None
        for g in reglas["reglas"]:
            if not desc.startswith(normalizar(g["prefijo"])):
                continue
            misma = True
            if _signo(g["signo"]) in ("±", signo):
                hit = g
                break
        if hit:
            x.update(concepto=hit["concepto"], destino=hit["destino"], origen="regla")
            tipo, _, clave = hit["destino"].partition(":")
            x["cuenta"] = reglas["cuentas"][clave]["codigo"] if tipo == "cuenta" else ""
        else:
            x.update(concepto=SIN_CONCEPTO, destino="revisar", cuenta="", origen="")
            x["avisos"].append("signo inesperado" if misma else "sin concepto: ninguna regla encaja")
        x.setdefault("tercero", "")
        x["factura_id"] = ""
        x.pop("sin_asiento", None)
        asg = x.get("asignada") or {}   # T9: lo que elegiste a mano manda y sobrevive a reclasificar
        if asg.get("cuenta"):
            x.update(destino="cuenta:manual", cuenta=asg["cuenta"], origen="manual")
        elif asg.get("factura_id"):
            x.update(destino="factura", cuenta="", origen="manual", factura_id=asg["factura_id"])
            usadas.add(asg["factura_id"])
        if x["destino"] == "revisar" and facturas and not asg.get("ninguna"):
            cs = candidatas(x, facturas, reglas["ventana_factura_dias"], usadas)
            if len(cs) == 1:
                x.update(destino="factura", cuenta="", origen="factura", factura_id=cs[0]["id"])
                x["avisos"].append(AV_CONFIRMAR)
                usadas.add(cs[0]["id"])
            elif cs:
                x["sin_asiento"] = True
                x["avisos"].append(AV_CUAL)
        if x["destino"] == "revisar" and not x.get("sin_asiento"):   # ADR-014 (FIX-2: con sin_asiento la memoria no toca): solo lo que iba a revisar; lo opaco nunca se recuerda
            ter = tercero_de(x.get("descripcion"), reglas)
            if ter:
                x["tercero"] = ter
                cod = _cuenta_recordada(memoria, ter)
                if cod:
                    x.update(destino="cuenta:recordada", cuenta=cod, origen="recordada")
        out.append(x)
    return out


def _fecha(v, n):
    if isinstance(v, datetime.datetime):
        return v.date().isoformat()
    if isinstance(v, datetime.date):
        return v.isoformat()
    raise ExtractoInvalido(f"fila {n}: la fecha no se entiende ({v!r})")


def _valor(v, n):
    if isinstance(v, bool) or not isinstance(v, (int, float)) or v != v or abs(v) > 1e12:
        raise ExtractoInvalido(f"fila {n}: el valor no es un numero ({v!r})")
    return int(round(v))


def leer_xlsx(datos: bytes) -> list:
    "Lee el extracto (hoja «Hoja 1», 4 columnas) y devuelve lineas con ocurrencia y clave; ExtractoInvalido(motivo) si algo falla."
    from openpyxl import load_workbook
    if not isinstance(datos, (bytes, bytearray)) or not datos:
        raise ExtractoInvalido("el archivo esta vacio")
    if len(datos) > MAX_BYTES:
        raise ExtractoInvalido("el archivo pesa mas de 5 MB")
    try:
        wb = load_workbook(io.BytesIO(bytes(datos)), read_only=True, data_only=True)
    except Exception as e:
        raise ExtractoInvalido("no es un Excel (.xlsx) valido") from e
    try:
        if HOJA not in wb.sheetnames:
            raise ExtractoInvalido(f"falta la hoja «{HOJA}»")
        filas = wb[HOJA].iter_rows(values_only=True)
        cab = next(filas, None)
        if cab is None or tuple(normalizar(c).strip() for c in cab[:4]) != COLUMNAS or any(c not in (None, "") for c in cab[4:]):
            raise ExtractoInvalido("las columnas deben ser Fecha, Descripción, Referencia y Valor")
        out, vistas = [], {}
        for n, f in enumerate(filas, 2):
            f = tuple(f[:4]) + (None,) * (4 - len(f))
            if all(c is None or str(c).strip() == "" for c in f):
                continue
            if len(out) >= MAX_FILAS:
                raise ExtractoInvalido(f"el archivo tiene mas de {MAX_FILAS} filas")
            fecha, valor = _fecha(f[0], n), _valor(f[3], n)
            desc = re.sub(r"\s+", " ", str(f[1] or "")).strip()
            if not desc:
                raise ExtractoInvalido(f"fila {n}: la descripcion esta vacia")
            ref = str(f[2] if f[2] is not None else "").strip()
            base = f"{fecha}|{normalizar(desc).strip()}|{ref}|{valor}"
            oc = vistas[base] = vistas.get(base, 0) + 1
            out.append({"clave": hashlib.sha1(f"{base}|{oc}".encode()).hexdigest()[:16], "fecha": fecha,
                        "descripcion": desc, "referencia": ref, "valor": valor, "ocurrencia": oc})
        return out
    except ExtractoInvalido:
        raise
    except Exception as e:
        raise ExtractoInvalido(f"no se pudo leer el archivo: {e}") from e
    finally:
        wb.close()
