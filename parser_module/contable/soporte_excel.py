"""Excel del documento soporte (R5): una hoja por documento con los campos del art. 4 de la Res. DIAN 167/2021
(número, fecha de generación, datos del vendedor, detalle de la operación, total e IVA) y la hoja RESUMEN.

Los valores salen de la copia fija guardada en `linea` (N2). La app no firma ni transmite: el Excel es para transcribir.
"""
from __future__ import annotations
import io
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from parser_module.contable.soporte import Documento

NEGRITA = Font(bold=True)
FONDO = PatternFill("solid", fgColor="DDEBF7")
COLS_RESUMEN = ["Número", "Proveedor", "Semana", "Total", "Fecha límite", "Estado"]
COLS_LINEA = ["Fecha de la compra", "Descripción", "Valor", "Papel de origen"]


def nombre_hoja(d: Documento) -> str:
    return f"{d.prefijo}-{d.numero}"[:31]


def libro(docs: list[Documento]) -> bytes:
    wb = Workbook()
    res = wb.active; res.title = "RESUMEN"
    res.append(COLS_RESUMEN)
    for d in docs:
        res.append([nombre_hoja(d), d.proveedor.nombre, d.semana, d.total, d.fecha_limite, d.estado])
        hoja = wb.create_sheet(nombre_hoja(d))
        for k, v in [("Número", nombre_hoja(d)), ("Fecha de generación", d.fecha_generacion),
                     ("Tipo de documento del vendedor", d.proveedor.tipo_id), ("Número de documento", d.proveedor.num_id),
                     ("Nombre del vendedor", d.proveedor.nombre), ("Dirección", d.proveedor.direccion or ""),
                     ("Municipio", d.proveedor.municipio or ""), ("Correo", d.proveedor.correo or ""),
                     ("Semana", d.semana), ("Estado", d.estado)]:
            hoja.append([k, v]); hoja.cell(hoja.max_row, 1).font = NEGRITA
        hoja.append([])
        hoja.append(COLS_LINEA)
        for c in hoja[hoja.max_row]: c.font = NEGRITA; c.fill = FONDO
        for l in d.lineas: hoja.append([l.fecha_compra, l.descripcion, l.valor, f"{l.origen_fecha} / {l.origen_doc_id}"])
        hoja.append(["", "IVA", d.iva]); hoja.append(["", "Total", d.total])
        for f in (hoja.max_row - 1, hoja.max_row): hoja.cell(f, 2).font = NEGRITA
        for col, w in zip("ABCD", (30, 40, 16, 28)): hoja.column_dimensions[col].width = w
    for c in res[1]: c.font = NEGRITA; c.fill = FONDO
    for col, w in zip("ABCDEF", (14, 32, 12, 14, 14, 14)): res.column_dimensions[col].width = w
    for row in res.iter_rows(min_row=2, min_col=4, max_col=4):
        for c in row: c.number_format = "#,##0"; c.alignment = Alignment(horizontal="right")
    buf = io.BytesIO(); wb.save(buf); return buf.getvalue()
