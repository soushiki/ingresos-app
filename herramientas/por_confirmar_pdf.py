"""T12 (M-27, R10, R4): lista para el contador de lo «por confirmar» en parser_module/contable/por_confirmar.json — las 8 decisiones
y los alias manual→PUC — con la regla que usa la app mientras tanto y cuántas filas del Excel manual (jun–oct) e ítems del libro afecta cada punto.
Contestar = editar por_confirmar.json («estado»: «confirmado», y «cuenta» o el alias si cambia) y luego «Reclasificar mes».
Uso: python herramientas/por_confirmar_pdf.py [--manual ruta.xlsx] [--libro datos/libro.json] [--salida docs/por-confirmar-contador.pdf]
Escribe el .html al lado del .pdf; el PDF lo imprime Edge o Chrome sin cabecera (si no hay ninguno, queda solo el .html)."""
from __future__ import annotations
import argparse, datetime as dt, html, json, re, subprocess, sys
from collections import Counter
from pathlib import Path

APP = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(APP)); sys.path.insert(0, str(APP / "herramientas"))
import minar_manual as mm
from parser_module.numeros import norm
from parser_module.contable import puc

POR_CONFIRMAR = APP / "parser_module" / "contable" / "por_confirmar.json"
LIBRO = APP / "datos" / "libro.json"
SALIDA = APP / "docs" / "por-confirmar-contador.pdf"
COMPRAS = ("factura", "gasto_caja", "caja_menor_manuscrito")
NAVEGADORES = [Path(r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe"), Path(r"C:\Program Files\Microsoft\Edge\Application\msedge.exe"),
               Path(r"C:\Program Files\Google\Chrome\Application\chrome.exe")]
e = lambda s: html.escape(str(s))
miles = lambda n: f"{n:,}".replace(",", ".")
pl = lambda n, uno, varios: f"<b>{miles(n)}</b> {uno if n == 1 else varios}"


def es_codigo(x: str) -> bool: return x.isdigit()


def toca(texto: str, palabras: list[str]) -> bool:
    t = norm(texto)
    return any(re.search(rf"\b{re.escape(norm(p))}\b", t) for p in palabras)


def nombres_por_codigo(ruta: Path) -> dict:
    """Columna C (nombre de la cuenta) de las hojas MOVIMIENTOS: {código: Counter(nombre)} — para la decisión de «un solo nombre»."""
    import openpyxl
    wb, out = openpyxl.load_workbook(ruta, read_only=True, data_only=True), {}
    for hoja in [n for n in wb.sheetnames if n.startswith("MOVIMIENTOS")]:
        for r in wb[hoja].iter_rows(min_row=9, max_col=3, values_only=True):
            if not isinstance(r[0], dt.datetime) or not r[1] or not r[2]: continue
            cod = str(r[1]).split(".")[0].strip()
            if cod.isdigit(): out.setdefault(cod, Counter())[" ".join(str(r[2]).split())] += 1
    return out


def items_libro(lib: dict) -> list[dict]:
    """Ítems de compra que cuentan (sin duplicados): {fecha, texto, cuenta, valor}. Papel viejo sin ítems = un ítem con su detalle."""
    out = []
    for fecha, dia in sorted(lib.items()):
        for d in dia.get("documentos", []):
            if d.get("tipo") not in COMPRAS or d.get("contar") is False: continue
            for it in d.get("items") or [dict(nombre=d.get("detalle", ""), cuenta=d.get("cuenta", ""), valor=d.get("valor", 0))]:
                out.append(dict(fecha=fecha, texto=f"{it.get('nombre', '')} {d.get('detalle', '')} {d.get('tercero', '')}", cuenta=str(it.get("cuenta") or ""),
                                valor=int(it.get("valor") or 0)))
    return out


def contar(afecta: list[str], filas: list[dict], items: list[dict]) -> dict:
    """Filas del manual y ítems del libro que tocan el punto: por código (codigo/cuenta igual) o por palabra (descripción, tercero)."""
    cods, pals = [a for a in afecta if es_codigo(a)], [a for a in afecta if not es_codigo(a)]
    gasto = lambda x: x["deb"] > 0 and not x["codigo"].startswith(mm.NO_GASTO)    # como minar: sin el lado de caja/banco del asiento
    fm = [x for x in filas if x["codigo"] in cods or (pals and gasto(x) and toca(f"{x['desc']} {x['tercero']}", pals))]
    fl = [x for x in items if x["cuenta"] in cods or (pals and toca(x["texto"], pals))]
    return dict(manual=len(fm), por_cuenta=Counter(x["codigo"] for x in fm).most_common(4), libro=len(fl), libro_valor=sum(x["valor"] for x in fl))


def nombre_puc(cod: str) -> str:
    try: return puc.nombre(cod) or ""
    except Exception: return ""


def construir_html(pc: dict, filas: list[dict], items: list[dict], nombres: dict, hoy: str) -> str:
    meses = sorted({x["mes"] for x in filas}); rango = f"{meses[0]} a {meses[-1]}" if meses else "sin datos"
    dias = sorted({x["fecha"] for x in items}); rango_l = f"{dias[0]} a {dias[-1]}" if dias else "sin días"
    pend = [d for d in pc["decisiones"] if d.get("estado") != "confirmado"]
    h = [f"""<!doctype html><html lang="es"><head><meta charset="utf-8"><title>Por confirmar con el contador</title><style>
@page{{size:A4;margin:12mm 13mm}}body{{font:9.6pt/1.38 "Segoe UI",Arial,sans-serif;color:#1d2733;margin:0}}
h1{{font-size:16.5pt;margin:0 0 2px}}h2{{font-size:12pt;margin:16px 0 5px;color:#17466f;border-bottom:1.5px solid #cfe0f0;padding-bottom:3px;page-break-after:avoid}}
.sub,.m{{color:#5b6b7b}}.m{{font-size:8.8pt}}table{{border-collapse:collapse;width:100%;margin:5px 0 7px;font-size:9pt;page-break-inside:avoid}}
th,td{{border:1px solid #d5dde6;padding:3px 6px;text-align:left;vertical-align:top}}th{{background:#eef4fa}}.n{{text-align:right;white-space:nowrap}}
.p{{border:1px solid #d5dde6;border-radius:6px;padding:6px 10px;margin:6px 0;page-break-inside:avoid}}.p b.q{{display:block;margin-bottom:3px}}
.op{{margin:2px 0}}.op:before{{content:"☐  ";font-size:11.5pt}}.linea{{border-bottom:1px solid #9aa8b6;height:16px;margin:3px 0}}
.est{{font-size:8.5pt;font-weight:700;padding:1px 7px;border-radius:9px;background:#fff1cc;color:#8a5d00;white-space:nowrap}}.ok{{background:#dff3e8;color:#1b6b46}}
</style></head><body>
<h1>Por confirmar con el contador · cuentas PUC</h1>
<p class="sub">Puntos donde la app tiene que elegir una cuenta y todavía no hay una respuesta suya. Mientras tanto usa la regla que sale en cada punto
(casi siempre: lo que más se usó en el Excel manual). Quedan <b>{len(pend)} de {len(pc['decisiones'])}</b> decisiones y
<b>{sum(1 for v in pc.get('alias_nota', {}).values() if v.get('estado') != 'confirmado')}</b> cambios de código por confirmar.</p>
<p class="m">«Manual» = filas del Excel contable de {e(rango)} ({miles(len(filas))} filas; por palabra solo cuenta el lado del gasto). «Libro» = ítems de compra guardados en la app, {e(rango_l)} (sin copias duplicadas).</p>
<h2>Decisiones</h2>"""]
    for d in pc["decisiones"]:
        c = contar(d.get("afecta", []), filas, items)
        cu = ", ".join(f"{k} ({v})" for k, v in c["por_cuenta"]) or "—"
        est = '<span class="est ok">confirmado</span>' if d.get("estado") == "confirmado" else '<span class="est">por confirmar</span>'
        h.append(f"""<div class="p"><b class="q">{d['id']}. {e(d['pregunta'])} {est}</b>
<table><tr><th style="width:24%">Hoy la app usa</th><td>{e(d.get('regla_usada', ''))}{f" → <b>{e(d['cuenta'])}</b>" if d.get('cuenta') else ""}</td></tr>
{f'<tr><th>Ya respondido</th><td>{e(d["respuesta"])}</td></tr>' if d.get("respuesta") else ""}
<tr><th>Afecta</th><td>Manual: {pl(c['manual'], 'fila', 'filas')} · en el manual fueron a {e(cu)}<br>Libro: {pl(c['libro'], 'ítem', 'ítems')}{f" · $ {miles(c['libro_valor'])}" if c['libro_valor'] else ""}</td></tr>
<tr><th>Busca</th><td class="m">{e(", ".join(d.get('afecta', [])))}</td></tr>""")
        if d["id"] == 8:
            for cod in d.get("afecta", []):
                ns = nombres.get(cod)
                if ns and len(ns) > 1: h.append(f"<tr><th>{e(cod)}</th><td>{e(' · '.join(f'{n} ({k})' for n, k in ns.most_common(5)))}</td></tr>")
        h.append("""</table><div class="op">De acuerdo con lo que usa la app.</div><div class="op">Otra cuenta / otra forma:</div><div class="linea"></div></div>""")
    h.append("""<h2>Cambios de código (alias manual → PUC)</h2>
<p>El Excel manual usa estos códigos; la app los cambia por la subcuenta PUC de la derecha. Así el contador no recibe códigos que no existen en el PUC.</p>
<table><tr><th>Código del manual</th><th class="n">Filas</th><th>Pasa a</th><th>Por qué</th><th>Estado</th><th>¿De acuerdo?</th></tr>""")
    for viejo, nuevo in pc.get("alias", {}).items():
        nota = pc.get("alias_nota", {}).get(viejo, {})
        n = sum(1 for x in filas if x["codigo"] == viejo)
        nom = next(iter(nombres.get(viejo, Counter()).most_common(1)), ("",))[0]
        h.append(f"""<tr><td>{e(viejo)}<br><span class="m">{e(nom)}</span></td><td class="n">{n}</td><td><b>{e(nuevo)}</b> <span class="m">{e(nombre_puc(nuevo))}</span></td>
<td>{e(nota.get('nota', ''))}</td><td>{'<span class="est ok">confirmado</span>' if nota.get('estado') == 'confirmado' else '<span class="est">por confirmar</span>'}</td><td>☐ Sí&nbsp; ☐ No: ______</td></tr>""")
    h.append(f"""</table>
<h2>Observaciones</h2><div class="p"><div class="linea"></div><div class="linea"></div><p class="m">Firma / fecha:</p><div class="linea"></div>
<p class="m" style="margin:6px 0 0">Hecho con la app de ingresos, {e(hoy)}. Para aplicar las respuestas: en <code>parser_module/contable/por_confirmar.json</code> cambiar «estado» a
«confirmado» (y «cuenta» o el alias si cambia); después «Reclasificar mes» en la pestaña ② Libro contable.</p></div></body></html>""")
    return "\n".join(h)


def imprimir_pdf(htm: Path, pdf: Path) -> bool:
    nav = next((p for p in NAVEGADORES if p.exists()), None)
    if not nav: return False
    subprocess.run([str(nav), "--headless=new", "--disable-gpu", "--no-pdf-header-footer", f"--print-to-pdf={pdf}", htm.resolve().as_uri()],
                   capture_output=True, timeout=120)
    return pdf.exists() and pdf.stat().st_mtime >= htm.stat().st_mtime - 1


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--manual", type=Path, default=mm.MANUAL); ap.add_argument("--libro", type=Path, default=LIBRO)
    ap.add_argument("--salida", type=Path, default=SALIDA)
    a = ap.parse_args(argv)
    pc = json.loads(POR_CONFIRMAR.read_text(encoding="utf8"))
    filas = mm.leer_filas(a.manual)
    lib = json.loads(a.libro.read_text(encoding="utf8")) if a.libro.exists() else {}
    htm = a.salida.with_suffix(".html")
    htm.write_text(construir_html(pc, filas, items_libro(lib), nombres_por_codigo(a.manual), dt.date.today().strftime("%d/%m/%Y")), encoding="utf8")
    ok = imprimir_pdf(htm, a.salida)
    print(f"{htm}" + (f"\n{a.salida}" if ok else "\n(no se imprimió el PDF: no hay Edge ni Chrome; abre el .html e imprímelo)"))


if __name__ == "__main__":
    main()
