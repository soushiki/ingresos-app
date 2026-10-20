"""Vocabulario de medio de pago (01c #8, T1): cada sinónimo → uno de los 8 medios; desconocido → vacío + aviso; nunca inventa."""
import json
from pathlib import Path
from parser_module.contable import documentos as D

fallos = []
def chk(cond, msg):
    print(("✓ " if cond else "✗ ") + msg)
    if not cond: fallos.append(msg)

CANON = {"Efectivo", "Transferencia", "Tarjeta", "QR", "Bre-B", "PSE", "App", "Crédito personal"}
chk(set(D.MEDIOS_CANON) == CANON and len(D.MEDIOS_CANON) == 8, "MEDIOS_CANON tiene los 8 medios")

# cada sinónimo del mapa (normalizacion.json) cae en un medio canónico
mapa = json.loads((Path(D.__file__).parent / "normalizacion.json").read_text(encoding="utf8"))["medio_canonico"]
chk(bool(mapa) and set(mapa.values()) <= CANON, "el mapa del json solo apunta a los 8")
for k, esperado in mapa.items():
    m, aviso = D.medio_canonico(k)
    chk(m == esperado and not aviso, f"«{k}» → {esperado} (dio {m!r})")

for txt, esp in [("Nequi", "Transferencia"), ("Daviplata", "Transferencia"), ("transf", "Transferencia"), ("Consignación", "Transferencia"),
                 ("Datáfono", "Tarjeta"), ("VISA", "Tarjeta"), ("débito", "Tarjeta"), ("Crédito tarjeta", "Tarjeta"),
                 ("Rappi", "App"), ("llave", "Bre-B"), ("Bre-B", "Bre-B"), ("Efectivo", "Efectivo"), ("QR", "QR"), ("PSE", "PSE"),
                 ("Crédito personal", "Crédito personal"), ("Tarjeta", "Tarjeta"), ("App", "App")]:
    chk(D.medio_canonico(txt) == (esp, ""), f"«{txt}» → {esp}")

for txt in ["(sin medio)", "", None, "Bitcoin", "Mixto", "xyz123"]:
    m, aviso = D.medio_canonico(txt)
    chk(m == "" and (aviso or txt in ("", None, "(sin medio)")), f"«{txt}» → vacío" + (" + aviso" if aviso else ""))
chk(D.medio_canonico("Bitcoin")[1] != "" and D.medio_canonico("Mixto")[1] != "", "desconocido y «Mixto» traen aviso")
chk(D.medio_canonico("(sin medio)")[0] == "" and D.medio_canonico("")[0] == "", "«(sin medio)» y vacío no inventan")

# _medio devuelve el canónico; el crudo queda en medio_detalle
chk(D._medio("Pago: Nequi") == "Transferencia", "_medio: Nequi → Transferencia")
chk(D._medio_crudo("Pago: Nequi") == "Nequi", "_medio_crudo conserva «Nequi»")
chk(D._medio("Visa débito") == "Tarjeta" and D._medio("nada que ver aquí") == "", "_medio: visa → Tarjeta; sin medio → vacío")
fac = D.interpretar("Factura de venta 123\nProveedor X SAS\nMétodo de pago: Daviplata\nTotal 50.000") if hasattr(D, "interpretar") else None
if fac:
    d = fac[0] if isinstance(fac, list) else fac
    chk(d["medio"] == "Transferencia" and d["medio_detalle"] == "Daviplata", f"factura: medio Transferencia, detalle Daviplata ({d['medio']}/{d.get('medio_detalle')})")

# T2: DATOS OCR y libro llevan el medio canónico; el crudo queda en medio_detalle; sin caja → medio vacío
from parser_module.contable import construir, leer_texto
def cmd(n, pago, prop="$"):
    return "\n".join(["SU RESTAURANTE", f"RECIBO #: {n}", "Fecha de Venta: 16/09/2026 9:00:00 p. m.", "SUBTOTAL CONSUMO $40,000", "**TOTAL CONSUMO>> $40000",
                      "*TOTAL A PAGAR> $40,000", f"PROPINA {prop}", "PAGADO $40,000", pago, "CAMBIO $"])
vs = [leer_texto(cmd(91001 + i, p))[0] for i, p in enumerate(["PAGA CON NEQUI $40,0", "PAGA CON TRANSFEREN", "PAGA CON VISA $40,000", "PAGA CON RAPPI $40,000"])]
sm = dict(vs[0], id="sinmedio", numero="91009", efectivo=0, tarjeta=0, transf=0, credito=0, avisos=[])
pr = dict(vs[0], id="conpropina", numero="91010", consumo=40000, propina=4000, efectivo=44000, tarjeta=0, transf=0, credito=0, avisos=[])
rr = construir("2026-09-16", vs + [sm, pr])
dd = rr["datos"]
chk(all(x["medio"] in CANON | {""} for x in dd), "DATOS OCR: todo medio es canónico o vacío (" + str({x["medio"] for x in dd}) + ")")
nq = next(x for x in dd if x["numero"] == "91001")
chk((nq["medio"], nq["medio_detalle"]) == ("Transferencia", "Nequi"), f"Nequi → Transferencia con detalle Nequi ({nq['medio']}/{nq['medio_detalle']})")
chk(next(x for x in dd if x["numero"] == "91003")["medio"] == "Tarjeta", "Datáfono/Visa → Tarjeta")
chk(next(x for x in dd if x["numero"] == "91004")["medio"] == "App", "Rappi → App")
x = next(x for x in dd if x["id"] == "sinmedio")
chk(x["medio"] == "" and "corrige el medio" in x["aviso"], "venta sin medio: medio vacío y aviso «corrige el medio»")
chk(any("corrige el medio" in a for a in rr["avisos"]), "el aviso de venta sin medio llega al libro")
chk(any(x["medio"] == "Efectivo" and x["propina"] == 4000 for x in dd), "la propina en efectivo viaja con la venta en Efectivo")
filas = rr["filas"]
chk(all(f["medio"] in CANON | {""} for f in filas), "libro: todo medio canónico o vacío")
chk(all(f["medio"] == "" for f in filas if f["codigo"] in ("13809501", "51959501") or f["descripcion"].startswith(("Recaudos con tarjeta, ", "Venta sin medio"))) and any(f["codigo"] == "13809501" for f in filas),
    "filas sin caja (puente, restaurante no efectivo) sin medio")
d_nq = [f for f in filas if f["codigo"] == "11200501" and f["debito"] and f["medio"] == "Transferencia"]
chk(len(d_nq) == 2 and {f["descripcion"].split(" · ")[0] for f in d_nq} == {"Recaudos a través de Transferencia Nequi", "Recaudos por transferencias"},
    f"Nequi no se junta con Transferencia: {[f['descripcion'] for f in d_nq]}")
chk(rr["diferencia"] == 0, "el día cuadra")
rec = [f for f in filas if f["codigo"] in ("11200501", "13452501") and f["debito"] and f["descripcion"].startswith(("Recaudos", "Ventas por"))]
chk(any(f["codigo"] == "13452501" for f in rec) and all("· Ordenes: 9100" in f["descripcion"] for f in rec),
    f"recaudos de banco y Rappi llevan «· Ordenes: …»: {[f['descripcion'] for f in rec]}")
rb =construir("2026-09-16", [], {"nequi": 1000, "transferencias": 2000, "llave": 500})
eb = {x["medio_detalle"]: (x["medio"], x["detalle"]) for x in rb["datos"] if x["tipo"] == "extracto"}
chk(eb.get("Nequi") == ("Transferencia", "Recaudos a través de Transferencia Nequi") and eb.get("Bre-B", ("",))[0] == "Bre-B", f"extracto: {eb}")

# T3: Mixto. Venta mixta sigue repartida (ninguna fila dice «Mixto»); factura de gasto Mixto toma el medio del mayor pago y conserva el aviso.
vm = leer_texto(cmd(91020, "PAGA CON VISA $15,000"))[0]
vm = dict(vm, id="mixta", numero="91020", consumo=40000, propina=0, efectivo=25000, tarjeta=15000, transf=0, credito=0, medio="Mixto", avisos=[])
fm = dict(vm, id="fmixta", tipo="factura", numero="F91", tercero="Proveedor X", detalle="Compra", valor=40000, cuenta="51959501", conf=1.0, motivo="", origen="x", medio="Mixto", efectivo=10000, tarjeta=30000)
fs = dict(fm, id="fsin", numero="F92", efectivo=0, tarjeta=0)
rm = construir("2026-09-16", [vm, fm, fs])
chk(not any("Mixto" == x["medio"] for x in rm["datos"]) and not any(f["medio"] == "Mixto" for f in rm["filas"]), "ninguna fila dice «Mixto»")
chk({x["medio"] for x in rm["datos"] if x["id"] == "mixta"} == {"Efectivo", "Tarjeta"}, "venta mixta sigue repartida en Efectivo y Tarjeta")
xf = next(x for x in rm["datos"] if x["id"] == "fmixta")
chk((xf["medio"], xf["medio_detalle"]) == ("Tarjeta", "Mixto"), f"factura Mixto toma el mayor pago ({xf['medio']}/{xf['medio_detalle']})")
chk(any("F91" in a and "revisa el medio" in a for a in rm["avisos"]), "factura Mixto mantiene el aviso «revisa el medio»")
chk(next(x for x in rm["datos"] if x["id"] == "fsin")["medio"] == "", "factura Mixto sin pagos leídos: medio vacío, sin inventar")

# T4 (R4): libro viejo (guardado antes de v1.5, con «Nequi», «Datáfono», «Mixto», «(sin medio)» en `medio`) → recalcular → Excel solo con los 8 o vacío
import io, copy
from openpyxl import load_workbook
from parser_module.contable import libro_a_excel
viejos = []
for n, med in [("92001", "Nequi"), ("92002", "Datáfono"), ("92003", "Mixto"), ("92004", "(sin medio)")]:
    v = dict(leer_texto(cmd(int(n), "PAGA CON NEQUI $40,0"))[0], id="v" + n, numero=n, medio=med, avisos=[])
    if med == "Mixto": v.update(efectivo=25000, tarjeta=15000, transf=0, credito=0)
    if med == "(sin medio)": v.update(efectivo=0, tarjeta=0, transf=0, credito=0)
    viejos.append(v)
rv = {"2026-09-17": construir("2026-09-17", copy.deepcopy(viejos))}
xl = load_workbook(io.BytesIO(libro_a_excel({k: r["filas"] for k, r in rv.items()}, "Neg", "1", {k: r["datos"] for k, r in rv.items()})))
for pref, fila_enc in (("DATOS OCR", 1), ("MOVIMIENTOS", 8)):
    ws = next((w for w in xl.worksheets if w.title.startswith(pref)), None)
    chk(ws is not None, f"Excel tiene la hoja {pref}")
    if ws is None: continue
    enc = [str(c.value or "").lower() for c in ws[fila_enc]]
    col = next(i for i, h in enumerate(enc) if "medio" in h and "detalle" not in h)
    vals = {r[col].value for r in ws.iter_rows(min_row=fila_enc + 1) if r[col].value is not None}
    chk(bool(vals) and vals <= CANON, f"libro viejo recalculado: {pref} solo con los 8 o vacío ({sorted(vals)})")
chk(not any(x["medio"] in ("Nequi", "Datáfono", "Mixto", "(sin medio)") for r in rv.values() for x in r["datos"]), "libro viejo: ningún Nequi/Datáfono/Mixto/(sin medio) en DATOS OCR")

print(f"\n{len(fallos)} fallos" if fallos else "\nTODO OK")
raise SystemExit(1 if fallos else 0)
