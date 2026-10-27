"""T4 · fitness del sitio del contador (ADR-027, ADR-028, S1, S4, S5): lock, lista blanca, cabeceras, sin URLs remotas,
y el núcleo del zip corre solo (intérprete limpio, sin el repo ni site-packages). Arma web/dist/ en una carpeta temporal
(baja las dependencias a web/cache/ la primera vez)."""
import io, json, pathlib, re, shutil, subprocess, sys, tempfile, time, zipfile
from openpyxl import load_workbook
sys.path.insert(0, str(pathlib.Path(__file__).parent / "herramientas"))
import web_build as wb   # noqa: E402
from parser_module.contable import libro   # noqa: E402

fallos = []
def chk(cond, msg):
    print(("✓ " if cond else "✗ ") + msg)
    if not cond: fallos.append(msg)

def rechaza(fn, *texto):
    "El build se niega con BuildError y el mensaje nombra la regla."
    try:
        fn()
    except wb.BuildError as e:
        return all(t in str(e) for t in texto)
    return False

AQUI = pathlib.Path(__file__).parent
CRUDO = (AQUI / "web" / "pruebas" / "datos" / "crudo_v1.json").read_bytes()
NEG, NIT, AHORA = "Restaurante de Prueba SAS", "900.123.456-7", "2026-10-26T09:30:00"
VER = "9.9.9"
tmp = pathlib.Path(tempfile.mkdtemp(prefix="probar_web_"))
try:
    # --- S5: lock fijado y verificado ---
    lock = wb.leer_lock()
    chk(lock["pyodide"]["version"] == "314.0.7" and {w["nombre"]: w["version"] for w in lock["wheels"]} == {"openpyxl": "3.1.5", "et_xmlfile": "2.0.0"},
        "S5: web.lock fija Pyodide 314.0.7, openpyxl 3.1.5, et_xmlfile 2.0.0")
    chk(all(len(d["sha256"]) == 64 and d["url"].startswith("https://") for d in wb.dependencias(lock)), "S5: cada dependencia trae url https y sha256")
    rutas = wb.bajar(lock)
    chk(all(wb.sha256(rutas[d["archivo"]].read_bytes()) == d["sha256"] for d in wb.dependencias(lock)), "S5: web/cache/ coincide con web.lock")
    cache_mala = tmp / "cache_mala"
    shutil.copytree(wb.CACHE, cache_mala)
    w = cache_mala / lock["wheels"][1]["archivo"]
    b = w.read_bytes()   # el último byte de un zip ya suele ser 0: se invierte uno del medio
    w.write_bytes(b[:100] + bytes([b[100] ^ 1]) + b[101:])
    chk(rechaza(lambda: wb.bajar(lock, cache_mala), "S5", "sha256", lock["wheels"][1]["archivo"]), "S5: un byte cambiado en un wheel detiene el build")

    # --- build completo en temporal ---
    dist = tmp / "dist"
    SIN_PUB = dict(pub=tmp / "sin_pub", plock=tmp / "sin_pub.lock")
    r = wb.armar(VER, dist=dist, src=tmp / "sin_src", **SIN_PUB)
    arch = r["archivos"]
    py = f"pyodide/{lock['pyodide']['version']}/"
    chk(sorted(a[len(py):] for a in arch if a.startswith(py)) == sorted(lock["pyodide"]["servir"]),
        f"S5: /{py} trae solo lo que carga el navegador ({len(lock['pyodide']['servir'])} archivos)")
    chk(sorted(a for a in arch if a.startswith("py/")) == sorted("py/" + w["archivo"] for w in lock["wheels"]), "S5: /py/ = los wheels del lock")
    chk(all(arch[f"py/{w['archivo']}"] == w["sha256"] for w in lock["wheels"]), "S5: los wheels servidos son los del lock")
    propios = [a for a in arch if not a.startswith(("py/", "pyodide/"))]
    chk(sorted(propios) == sorted(["_headers", "index.html", f"v{VER}/manifiesto.json", f"v{VER}/nucleo.zip"]), f"ADR-027: dist = raíz + /v{VER}/: {propios}")

    # --- ADR-028: lista blanca exacta ---
    nz = (dist / f"v{VER}/nucleo.zip").read_bytes()
    with zipfile.ZipFile(io.BytesIO(nz)) as z:
        nombres = z.namelist()
        vacios = {n: z.read(n) for n in wb.VACIOS}
    chk(sorted(nombres) == sorted(wb.LISTA_BLANCA), "ADR-028: nucleo.zip = exactamente la lista blanca")
    chk(all(v == b"" for v in vacios.values()), "ADR-028: los dos __init__.py van vacíos (sin clasificador ni parser OCR)")
    chk(not any(wb.PROHIBIDO.search(n) for n in nombres), "S1: ningún nombre prohibido en nucleo.zip")
    chk(wb.nucleo_zip() == nz, "ADR-027: nucleo.zip determinista (mismo código → mismos bytes)")
    man = json.loads((dist / f"v{VER}/manifiesto.json").read_text(encoding="utf8"))
    chk(man == {"version": VER, "pyodide": "314.0.7", "wheels": [w["archivo"] for w in lock["wheels"]], "nucleo_sha256": wb.sha256(nz)},
        "ADR-027: manifiesto con versión, Pyodide, wheels y sha256 del núcleo")
    for malo in ["parser_module/contable/terceros_cuentas.json", "parser_module/contable/modelo_manual.json", "datos/libro.json",
                 "parser_module/contable/contable.json", "datos/cierres.db"]:
        chk(rechaza(lambda: wb.nucleo_zip(lista=wb.LISTA_BLANCA + [malo]), "ADR-028", malo), f"S1: {malo} en la lista detiene el build")

    # --- web/src: nombres prohibidos y salidas ---
    casos = {"a.js": 'fetch("https://evil.example/x")', "b.js": "importScripts('http://cdn/x.js')", "c.js": "el.innerHTML = nombre",
             "d.js": "eval (codigo)", "e.js": "new Function('return 1')", "f.js": "py.loadPackage('micropip')",
             "g.html": '<script src="https://cdn.jsdelivr.net/pyodide.js"></script>'}
    for nombre, texto in casos.items():
        s = tmp / f"src_{nombre}"
        s.mkdir()
        (s / nombre).write_text(texto, encoding="utf8")
        chk(rechaza(lambda: wb.armar(VER, dist=tmp / "dist_x", src=s), "ADR-028", nombre), f"ADR-028: web/src/{nombre} «{texto[:30]}» detiene el build")
    s = tmp / "src_terceros"
    (s / "datos").mkdir(parents=True)
    (s / "datos" / "x.json").write_text("{}", encoding="utf8")
    chk(rechaza(lambda: wb.armar(VER, dist=tmp / "dist_x", src=s), "S1", "datos/x.json"), "S1: web/src/datos/ detiene el build")
    s = tmp / "src_ok"
    s.mkdir()
    (s / "app.js").write_text('import { loadPyodide } from "/pyodide/314.0.7/pyodide.mjs";\nel.textContent = "hola";\n', encoding="utf8")
    r2 = wb.armar(VER, dist=tmp / "dist_ok", src=s)
    chk(f"v{VER}/app.js" in r2["archivos"], "web/src/ limpio se copia a /vX.Y.Z/")
    chk(rechaza(lambda: wb.armar("1.10", dist=tmp / "dist_x"), "ADR-027"), "ADR-027: versión que no es X.Y.Z detiene el build")

    # --- sin URLs remotas ni ejecución dinámica en lo propio del sitio ---
    sucios = [a for a in propios if pathlib.Path(a).suffix in wb.CODIGO and wb.SALIDA.search((dist / a).read_text(encoding="utf8"))]
    chk(not sucios, f"S3: sin http(s)://, micropip, innerHTML, eval ni new Function en lo propio: {sucios}")
    with zipfile.ZipFile(io.BytesIO(nz)) as z:
        sucios = [n for n in z.namelist() if n.endswith(".py") and wb.SALIDA.search(z.read(n).decode("utf8"))]
    chk(not sucios, f"S3: el núcleo Python no trae URLs ni eval: {sucios}")

    # --- S4: cabeceras ---
    S4 = ("default-src 'none'; script-src 'self' 'wasm-unsafe-eval'; worker-src 'self' blob:; connect-src 'self'; img-src 'self' blob: data:; "
          "style-src 'self'; frame-ancestors 'none'; base-uri 'none'; form-action 'none'")
    bloques, actual = {}, None
    for linea in (dist / "_headers").read_text(encoding="utf8").splitlines():
        if linea and not linea[0].isspace():
            actual = bloques.setdefault(linea.strip(), {})
        elif linea.strip():
            n, _, v = linea.strip().partition(":")
            actual[n.strip()] = v.strip()
    raiz = bloques.get("/*", {})
    chk(raiz.get("Content-Security-Policy") == S4, "S4: la CSP de /* es exactamente la de S4")
    spike = (AQUI / "spikes" / "pyodide" / "csp" / "_headers").read_text(encoding="utf8")
    chk(f"Content-Security-Policy: {S4}" in spike, "S4: la misma CSP que pasó el spike T1 en Chrome/Edge")
    chk(raiz.get("X-Content-Type-Options") == "nosniff" and raiz.get("Referrer-Policy") == "no-referrer" and raiz.get("Permissions-Policy") == "",
        "S4: nosniff, no-referrer, Permissions-Policy vacía")
    chk("Cache-Control" not in raiz, "ADR-027: la raíz sin caché larga (Netlify junta reglas: Cache-Control nunca en /*)")
    rutas_inm = {k for k, v in bloques.items() if v.get("Cache-Control") == wb.INMUTABLE}
    chk(rutas_inm == {"/pyodide/314.0.7/*", "/py/*", f"/v{VER}/*"}, f"ADR-027: immutable solo en rutas versionadas: {sorted(rutas_inm)}")
    idx = (dist / "index.html").read_text(encoding="utf8")
    chk(f'href="v{VER}/"' in idx and "<script" not in idx and "style=" not in idx, "ADR-027: índice raíz lista la versión, sin script ni estilo en línea")

    # --- T10 · ADR-027/R6: versiones publicadas inmutables ---
    chk(isinstance(wb.verificar_publicadas(), dict), f"ADR-027: web/publicadas/ del repo = publicadas.lock ({len(wb.leer_publicadas())} versiones)")
    VIEJA, pub, plock = "9.9.8", tmp / "pub", tmp / "pub.lock"
    T10 = dict(pub=pub, plock=plock)
    s = tmp / "src_v1"
    s.mkdir()
    (s / "app.js").write_text("el.textContent = 'v vieja';\n", encoding="utf8")
    wb.armar(VIEJA, dist=tmp / "dist_v1", src=s, **T10)
    sellos = wb.publicar(VIEJA, dist=tmp / "dist_v1", **T10)
    chk(sorted(sellos) == ["app.js", "manifiesto.json", "nucleo.zip"] and wb.leer_publicadas(plock) == {VIEJA: sellos},
        f"ADR-027: --publicar copia /v{VIEJA}/ y guarda su sha256 en el lock")
    chk(rechaza(lambda: wb.publicar(VIEJA, dist=tmp / "dist_v1", **T10), "ADR-027", "ya está publicada"), "ADR-027: publicar la misma versión dos veces se niega")
    (s / "app.js").write_text("el.textContent = 'v nueva';\n", encoding="utf8")
    d2 = tmp / "dist_v2"
    r3 = wb.armar(VER, dist=d2, src=s, **T10)
    chk(all(r3["archivos"][f"v{VIEJA}/{k}"] == h for k, h in sellos.items()) and (d2 / f"v{VER}/app.js").read_text(encoding="utf8").endswith("nueva';\n"),
        f"ADR-027: el build trae v{VIEJA} intacta (web/src nuevo solo entra en v{VER})")
    idx = (d2 / "index.html").read_text(encoding="utf8")
    chk(r3["versiones"] == [VER, VIEJA] and idx.index(f'href="v{VER}/"') < idx.index(f'href="v{VIEJA}/"'), "ADR-027: el índice lista todas, la última arriba")
    chk(f"/v{VIEJA}/*\n  Cache-Control: {wb.INMUTABLE}" in (d2 / "_headers").read_text(encoding="utf8"), f"ADR-027: /v{VIEJA}/* también immutable")
    r4 = wb.armar(VIEJA, dist=tmp / "dist_v3", src=s, **T10)
    chk(r4["publicada"] and all(r4["archivos"][f"v{VIEJA}/{k}"] == h for k, h in sellos.items()), f"ADR-027: armar v{VIEJA} otra vez usa la copia publicada, no web/src")
    nz_pub = pub / f"v{VIEJA}" / "nucleo.zip"
    original = nz_pub.read_bytes()
    nz_pub.write_bytes(original[:50] + bytes([original[50] ^ 1]) + original[51:])
    chk(rechaza(lambda: wb.armar(VER, dist=tmp / "dist_x", src=s, **T10), "ADR-027", f"v{VIEJA} cambió", "nucleo.zip", "tag nuevo"),
        "ADR-027: cambiar 1 byte de una versión publicada detiene el build")
    nz_pub.write_bytes(original)
    (pub / f"v{VIEJA}" / "extra.js").write_text("x", encoding="utf8")
    chk(rechaza(lambda: wb.armar(VER, dist=tmp / "dist_x", src=s, **T10), "ADR-027", "extra.js"), "ADR-027: un archivo de más en una publicada detiene el build")
    (pub / f"v{VIEJA}" / "extra.js").unlink()
    (pub / f"v{VIEJA}" / "app.js").rename(tmp / "app.js.bak")
    chk(rechaza(lambda: wb.armar(VER, dist=tmp / "dist_x", src=s, **T10), "ADR-027", "app.js"), "ADR-027: un archivo borrado de una publicada detiene el build")
    (tmp / "app.js.bak").rename(pub / f"v{VIEJA}" / "app.js")
    shutil.copytree(pub / f"v{VIEJA}", pub / "v9.9.7")
    chk(rechaza(lambda: wb.armar(VER, dist=tmp / "dist_x", src=s, **T10), "ADR-027", "v9.9.7", "no está en publicadas.lock"), "ADR-027: una carpeta copiada a mano detiene el build")
    shutil.rmtree(pub / "v9.9.7")
    otro = json.loads(json.dumps(lock)); otro["pyodide"]["version"] = "0.0.1"
    chk(rechaza(lambda: wb.verificar_publicadas(pub, plock, otro), "ADR-027", "Pyodide 314.0.7"), "ADR-027: web.lock sin el Pyodide de una publicada detiene el build")
    chk(isinstance(wb.verificar_publicadas(pub, plock), dict), "ADR-027: restaurada, la publicada vuelve a pasar")

    # --- ADR-028: el núcleo del zip corre solo (como en Pyodide: zip desempacado + wheels, sin repo ni site-packages) ---
    nucleo_dir = tmp / "nucleo"
    with zipfile.ZipFile(io.BytesIO(nz)) as z:
        z.extractall(nucleo_dir)
    (tmp / "crudo.json").write_bytes(CRUDO)
    codigo = ("import sys, json; sys.path[:0] = json.loads(sys.argv[1])\n"
              "from parser_module.contable import libro\n"
              "assert libro.__file__.startswith(sys.argv[2]), libro.__file__\n"
              "import openpyxl; assert '.whl' in openpyxl.__file__, openpyxl.__file__\n"
              f"g = libro.generar(open(sys.argv[3], 'rb').read(), None, None, {NEG!r}, {NIT!r}, {AHORA!r})\n"
              "open(sys.argv[4], 'wb').write(g['xlsx'])\n"
              "print(json.dumps(sorted(m for m in sys.modules if m.startswith('parser_module'))))")
    ruta_py = [str(nucleo_dir)] + [str(dist / "py" / w["archivo"]) for w in lock["wheels"]]
    p = subprocess.run([sys.executable, "-I", "-S", "-c", codigo, json.dumps(ruta_py), str(nucleo_dir), str(tmp / "crudo.json"), str(tmp / "zip.xlsx")],
                       capture_output=True, text=True, timeout=120)
    chk(p.returncode == 0, "ADR-028: libro.generar corre con solo nucleo.zip + wheels (intérprete limpio)" + ("" if p.returncode == 0 else f"\n{p.stderr[-800:]}"))
    if p.returncode == 0:
        cargados = json.loads(p.stdout.strip().splitlines()[-1])
        chk(not any(m.endswith(("clasificar", "parser", "soporte", "soporte_db")) for m in cargados), f"ADR-028: no carga el clasificador ni el parser OCR ({len(cargados)} módulos)")
        def valores(b):
            libro_x = load_workbook(io.BytesIO(b))
            return {(ws.title, c.coordinate): (c.value, c.number_format) for ws in libro_x for f in ws.iter_rows() for c in f if c.value is not None}
        ref = libro.generar(CRUDO, None, None, NEG, NIT, AHORA)["xlsx"]
        a, b = valores((tmp / "zip.xlsx").read_bytes()), valores(ref)
        chk(a == b, f"R1: el Excel del zip = el del repo ({len(a)} celdas, {sum(a.get(k) != b.get(k) for k in set(a) | set(b))} distintas)")

    # --- T9 · S2/S4: archivos de ataque (web/pruebas/ataques/) → mensaje en español, rápido, sin traza; la página no los ejecuta ---
    sys.path.insert(0, str(AQUI / "web" / "pruebas" / "ataques"))
    from ataques import CASOS, NOMBRE_SCRIPT   # noqa: E402
    for nombre, (datos, esperado, _) in CASOS.items():
        t0 = time.perf_counter()
        try:
            libro.generar(datos, None, None, NEG, NIT, AHORA); m = "SIN ERROR"
        except libro.EntradaInvalida as e: m = str(e)
        except Exception as e: m = f"OTRA EXCEPCIÓN {type(e).__name__}"   # noqa: BLE001
        seg = time.perf_counter() - t0
        chk(esperado in m and seg < 2 and not any(x in m for x in ("Traceback", ".py", "\\")), f"S2 {nombre} ({len(datos):,} B) → «{m}» en {seg:.2f} s")
    js = (AQUI / "web" / "src" / "app.js").read_text(encoding="utf8")
    usos = [l.strip() for l in js.splitlines() if re.search(r"\bfe?\.name\b", l)]
    chk(usos and all("textContent" in l or l.startswith("return error(") for l in usos), f"S2: el nombre del archivo solo entra con textContent ({len(usos)} usos)")
    chk(all("name" not in l for l in js.splitlines() if "postMessage" in l), "S2: el nombre del archivo no viaja al Worker")
    envio = js[js.index('form.addEventListener("submit"'):]
    chk(envio.index("f.size > TOPE_CRUDO") < envio.index("arrayBuffer()"), "S2: el tope de 20 MB se mira antes de leer el archivo (25 MB no llega a memoria)")
    wk = (AQUI / "web" / "src" / "motor.worker.js").read_text(encoding="utf8")
    topes = {k: int(n) * 2 ** 20 for k, n in re.findall(r"(crudo|extracto|memoria): \[(\d+) \* MB", wk)}
    chk(topes == {"crudo": libro.MAX_CRUDO, "extracto": libro.MAX_EXTRACTO, "memoria": libro.MAX_MEMORIA}, f"S2: los topes del Worker = los de libro.generar {topes}")
    chk("<" in NOMBRE_SCRIPT and (AQUI / "web" / "pruebas" / "ataques" / "CHECKLIST.md").exists(), "S4: checklist en navegador (nombre con <script>) en web/pruebas/ataques/CHECKLIST.md")
finally:
    shutil.rmtree(tmp, ignore_errors=True)

print(f"\n{'OK' if not fallos else f'{len(fallos)} FALLOS'}")
sys.exit(1 if fallos else 0)
